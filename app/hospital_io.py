"""Hospital data in and out of Excel, and clearing the demo data before going live.

One workbook holds one sheet per kind of record (rooms, beds, equipment types and
units, staff, supplies, procedures, preference cards). Download the current data
(or a blank form), edit it in Excel, and upload it back (.xlsx, or one sheet as CSV).

An upload is always checked first: every row goes through the same save code as the
hospital data pages, inside one transaction that is then rolled back, and the
result lists what would be added, changed or is wrong (with the Excel row number).
Only a file without errors can be imported; it is saved in one transaction.

Rows are matched to existing records by code (rooms, beds, staff, equipment serials,
supplies), by name (equipment types; procedures without an ICD-9-CM code) or by
surgeon + procedure (preference cards). Records that are not in the file are kept.
A column that is left out keeps the current values; a blank cell clears optional
fields and keeps the current value of numbers and choices.
"""
from __future__ import annotations

import csv
import io
import itertools
import re
import zipfile
from dataclasses import dataclass
from datetime import datetime, time, timedelta

import psycopg

from .admin_store import (
    BED_STATUSES,
    BED_TYPES,
    CARD_ROLES,
    CHECK_DAYS,
    EQUIPMENT_CATEGORIES,
    INFECTION_LEVELS,
    ROLES,
    ROOM_STATUSES,
    ROOM_TYPES,
    SUPPLY_CATEGORIES,
    HospitalAdmin,
)
from .seed import HISTORY_TAG
from .store import StoreError

MAX_BYTES = 5 * 1024 * 1024  # uploaded file
MAX_SHEETS = 20
MAX_UNZIPPED = 60 * 1024 * 1024  # an .xlsx is a zip: refuse ones that expand too far
MAX_ROWS = 5000  # data rows per sheet
FORM_ROWS = 500  # rows given dropdowns / text format in the downloaded form
GUIDE = "วิธีกรอก"
YES = {"ใช่", "yes", "y", "true", "1", "✓", "/"}
NO = {"ไม่", "ไม่ใช่", "no", "n", "false", "0"}
WIPE_PHRASE = "ล้างข้อมูล"
DEMO_NAME = "ผู้ป่วยตัวอย่าง %"  # patient names of the demo day's cases (app/db_store.py seed_demo)


@dataclass(frozen=True)
class Col:
    key: str
    label: str
    kind: str = "text"  # text | int | float | bool | enum | list | time | skills | counts | info
    required: bool = False  # must be filled for a new record
    choices: tuple = ()
    aliases: tuple = ()  # (Thai word, code)
    keep_blank: bool = False  # a blank cell keeps the current value (numbers and choices)
    help: str = ""
    example: str = ""
    width: int = 16
    download: bool = True  # filled in the downloaded file (False: left blank to type into)


@dataclass(frozen=True)
class Sheet:
    kind: str
    title: str
    cols: tuple
    help: str
    needs: tuple | None = None  # (table, column) the database must have


def _yes_no(word: bool) -> str:
    return "ใช่" if word else "ไม่"


ROOMS = Sheet("rooms", "ห้องผ่าตัด", (
    Col("code", "รหัสห้อง", required=True, help="1-20 ตัว (A-Z 0-9 - _ .)", example="OR-1", width=12),
    Col("type", "ประเภทห้อง", "enum", choices=ROOM_TYPES, keep_blank=True,
        aliases=(("มาตรฐาน", "STANDARD"), ("ห้องมาตรฐาน", "STANDARD"), ("ความดันลบ", "NEGATIVE_PRESSURE"),
                 ("ห้องความดันลบ", "NEGATIVE_PRESSURE"), ("ห้องhybrid", "HYBRID")),
        help="STANDARD = มาตรฐาน · HYBRID · NEGATIVE_PRESSURE = ความดันลบ", example="STANDARD", width=20),
    Col("capabilities", "ความสามารถของห้อง", "list",
        help="คั่นด้วยจุลภาค เช่น LAMINAR_FLOW, LEAD_LINED_XRAY, NEGATIVE_PRESSURE, ROBOTIC_READY",
        example="LAMINAR_FLOW, ROBOTIC_READY", width=32),
    Col("base_turnover_mins", "ทำความสะอาดปกติ (นาที)", "int", keep_blank=True, example="30", width=14),
    Col("infectious_turnover_mins", "ทำความสะอาดหลังเคสติดเชื้อ (นาที)", "int", keep_blank=True, example="60", width=16),
    Col("status", "สถานะ", "enum", choices=ROOM_STATUSES, keep_blank=True,
        aliases=(("ใช้งานได้", "AVAILABLE"), ("ซ่อมบำรุง", "MAINTENANCE"), ("ปิดใช้งาน", "CLOSED")),
        help="AVAILABLE = ใช้งานได้ · MAINTENANCE = ซ่อมบำรุง · CLOSED = ปิดใช้งาน", example="AVAILABLE", width=14),
    Col("elective_start", "เริ่มรับเคส elective", "time", help="HH:MM · ว่าง = ไม่จำกัด", example="08:00", width=12),
    Col("elective_end", "ปิดรับเคส elective", "time", help="HH:MM · ว่าง = ไม่จำกัด", example="16:00", width=12),
), "หนึ่งแถวต่อหนึ่งห้อง · จับคู่กับห้องเดิมด้วยรหัสห้อง")

BEDS = Sheet("beds", "เตียงพักฟื้น", (
    Col("code", "รหัสเตียง", required=True, help="1-20 ตัว (A-Z 0-9 - _ .)", example="PACU-1", width=12),
    Col("type", "ประเภทเตียง", "enum", choices=BED_TYPES, keep_blank=True,
        aliases=(("ห้องพักฟื้น", "PACU"), ("icu", "ICU_SURGICAL"), ("icuศัลยกรรม", "ICU_SURGICAL")),
        help="PACU = ห้องพักฟื้น · ICU_SURGICAL = ICU ศัลยกรรม", example="PACU", width=16),
    Col("status", "สถานะ", "enum", choices=BED_STATUSES, keep_blank=True,
        aliases=(("ใช้งานได้", "AVAILABLE"), ("เปิด", "AVAILABLE"), ("ปิด", "CLOSED"), ("ปิดใช้งาน", "CLOSED")),
        help="AVAILABLE = ใช้งานได้ · CLOSED = ปิด", example="AVAILABLE", width=14),
    Col("note", "หมายเหตุ", width=30),
), "หนึ่งแถวต่อหนึ่งเตียง · จับคู่ด้วยรหัสเตียง")

TYPES = Sheet("types", "ชนิดเครื่องมือ", (
    Col("name", "ชื่อชนิดเครื่องมือ", required=True, example="LAPAROSCOPY_TRAY", width=30),
    Col("category", "หมวด", "enum", choices=EQUIPMENT_CATEGORIES, keep_blank=True,
        aliases=(("เครื่องขนาดใหญ่", "HEAVY_MACHINE"), ("ชุดเครื่องมือ", "INSTRUMENT_TRAY"), ("อื่นๆ", "OTHER")),
        help="HEAVY_MACHINE = เครื่องขนาดใหญ่ · INSTRUMENT_TRAY = ชุดเครื่องมือ · OTHER = อื่น ๆ",
        example="INSTRUMENT_TRAY", width=18),
    Col("requires_sterilization", "ต้องอบฆ่าเชื้อ", "bool", keep_blank=True, help="ใช่ / ไม่", example="ใช่", width=12),
    Col("sterilization_mins", "เวลาอบฆ่าเชื้อ (นาที)", "int", keep_blank=True, example="120", width=14),
), "หนึ่งแถวต่อหนึ่งชนิด · จับคู่ด้วยชื่อ (ไม่สนตัวพิมพ์เล็ก/ใหญ่)")

UNITS = Sheet("units", "เครื่องมือรายชิ้น", (
    Col("serial", "หมายเลขเครื่อง", required=True, example="LAP-01", width=16),
    Col("type", "ชนิดเครื่องมือ", required=True, help="ชื่อตามแผ่น ชนิดเครื่องมือ", example="LAPAROSCOPY_TRAY", width=30),
    Col("location", "ที่เก็บ", example="CSSD", width=16),
), "หนึ่งแถวต่อหนึ่งเครื่อง/ชุด · จับคู่ด้วยหมายเลขเครื่อง · สถานะรายวันอัปเดตที่แท็บเครื่องมือของ CSSD")

STAFF = Sheet("staff", "บุคลากร", (
    Col("code", "รหัสพนักงาน", required=True, example="D-0012", width=14),
    Col("name", "ชื่อ-นามสกุล", required=True, example="นพ.สมชาย ใจดี", width=28),
    Col("role", "บทบาท", "enum", required=True, choices=ROLES, keep_blank=True,
        aliases=(("ศัลยแพทย์", "SURGEON"), ("วิสัญญีแพทย์", "ANESTHESIOLOGIST"), ("วิสัญญี", "ANESTHESIOLOGIST"),
                 ("พยาบาลส่งเครื่องมือ", "SCRUB_NURSE"), ("scrub", "SCRUB_NURSE"),
                 ("พยาบาลช่วยรอบนอก", "CIRCULATING_NURSE"), ("circulate", "CIRCULATING_NURSE")),
        help="SURGEON = ศัลยแพทย์ · ANESTHESIOLOGIST = วิสัญญีแพทย์ · SCRUB_NURSE = พยาบาลส่งเครื่องมือ · "
             "CIRCULATING_NURSE = พยาบาลช่วยรอบนอก", example="SURGEON", width=20),
    Col("specialty", "ความเชี่ยวชาญ", example="General Surgery", width=20),
    Col("skills", "ทักษะ", "list", help="คั่นด้วยจุลภาค ใช้ A-Z 0-9 _ เช่น ROBOTIC_CERTIFIED", width=26),
    Col("max_hours_per_week", "ชั่วโมงสูงสุดต่อ 7 วัน", "float", keep_blank=True, example="48", width=14),
    Col("max_hours_per_day", "ชั่วโมงสูงสุดต่อวัน", "float", keep_blank=True, example="10", width=14),
    Col("min_rest_hours", "พักหลังเคสดึก (ชม.)", "float", keep_blank=True, example="8", width=14),
    Col("active", "ใช้งาน", "bool", keep_blank=True, help="ไม่ = ปิดใช้งาน (แทนการลบ)", example="ใช่", width=10),
), "หนึ่งแถวต่อหนึ่งคน · จับคู่ด้วยรหัสพนักงาน · ตารางเวรใช้ไฟล์ Excel แยก (แท็บเวรและวันลา)")

SUPPLIES = Sheet("supplies", "วัสดุและรากเทียม", (
    Col("code", "รหัสวัสดุ", required=True, help="1-50 ตัว (A-Z 0-9 - _ .)", example="KNEE-IMP-5", width=16),
    Col("name", "ชื่อวัสดุ", required=True, example="ข้อเข่าเทียม ขนาด 5", width=30),
    Col("category", "หมวด", "enum", choices=SUPPLY_CATEGORIES, keep_blank=True,
        aliases=(("รากเทียม", "IMPLANT"), ("วัสดุสิ้นเปลือง", "CONSUMABLE"), ("สิ้นเปลือง", "CONSUMABLE")),
        help="IMPLANT = รากเทียม · CONSUMABLE = วัสดุสิ้นเปลือง", example="IMPLANT", width=14),
    Col("unit", "หน่วย", keep_blank=True, example="ชิ้น", width=10),
    Col("on_hand_now", "คงคลังในระบบ (ดูอย่างเดียว)", "info", help="ยอดตอนดาวน์โหลด ไม่ใช้ตอนนำเข้า", width=14),
    Col("on_hand", "นับได้จริง (กรอกเมื่อนับสต็อก)", "int", keep_blank=True, download=False,
        help="ว่าง = ไม่เปลี่ยนยอด · กรอกแล้วบันทึกเป็นการนับสต็อก (ยอดในระบบเปลี่ยนเองเมื่อผ่าตัดเสร็จหรือรับของ "
             "จึงไม่ใส่ยอดเก่าให้) · รายการใหม่: ยอดเริ่มต้น", example="4", width=16),
    Col("reorder_level", "จุดสั่งซื้อ", "int", keep_blank=True, example="2", width=12),
    Col("active", "ใช้งาน", "bool", keep_blank=True, help="ไม่ = ปิดใช้งาน", example="ใช่", width=10),
), "หนึ่งแถวต่อหนึ่งรายการ · จับคู่ด้วยรหัสวัสดุ", needs=("supply_items", "item_id"))

PROCEDURES = Sheet("procedures", "หัตถการ", (
    Col("code", "รหัส ICD-9-CM", help="ไม่บังคับ · ถ้ามีจะใช้จับคู่กับหัตถการเดิม", example="51.23", width=14),
    Col("name", "ชื่อหัตถการ", required=True, example="Laparoscopic cholecystectomy", width=34),
    Col("duration", "เวลาผ่าตัดมาตรฐาน (นาที)", "int", keep_blank=True, example="90", width=14),
    Col("infection", "ความเสี่ยงติดเชื้อ", "enum", choices=INFECTION_LEVELS, keep_blank=True,
        aliases=(("สะอาด", "CLEAN"), ("ปนเปื้อน", "CONTAMINATED"), ("ติดเชื้อ", "DIRTY")),
        help="CLEAN = สะอาด · CONTAMINATED = ปนเปื้อน · DIRTY = ติดเชื้อ", example="CLEAN", width=16),
    Col("icu", "ต้องเข้า ICU หลังผ่าตัด", "bool", keep_blank=True, help="ใช่ / ไม่", example="ไม่", width=12),
    Col("blood", "ต้องจองเลือด", "bool", keep_blank=True, help="ใช่ = รายการเช็กความพร้อมถามเรื่องจองเลือด",
        example="ไม่", width=12),
    Col("recovery_mins", "เวลาพักฟื้น (นาที)", "int", help="ว่าง = มาตรฐาน (PACU 60 นาที / ICU 24 ชม.)", width=14),
    Col("caps", "ห้องต้องมี", "list", help="ความสามารถของห้อง คั่นด้วยจุลภาค", example="LAMINAR_FLOW", width=22),
    Col("skills", "ทักษะที่ทีมต้องมี", "skills",
        help="บทบาท: ทักษะ · คั่นแต่ละบทบาทด้วย ; เช่น SCRUB_NURSE: ROBOTIC_CERTIFIED", width=34),
    Col("supplies", "วัสดุที่ใช้", "counts",
        help="รหัสวัสดุ x จำนวน · คั่นด้วย ; เช่น KNEE-IMP-5 x1; BONE-CEMENT x2", width=30),
), "หนึ่งแถวต่อหนึ่งหัตถการ · จับคู่ด้วยรหัส ICD-9-CM หรือชื่อ")

CARDS = Sheet("cards", "Preference card", (
    Col("surgeon_code", "รหัสศัลยแพทย์", required=True, help="รหัสพนักงานตามแผ่นบุคลากร", example="D-0012", width=14),
    Col("surgeon", "ชื่อศัลยแพทย์ (ไม่ต้องกรอก)", "info", help="ระบบใส่ให้ตอนดาวน์โหลด ไม่ใช้ตอนนำเข้า", width=24),
    Col("procedure", "หัตถการ", required=True, help="ชื่อหรือรหัส ICD-9-CM ตามแผ่นหัตถการ",
        example="Laparoscopic cholecystectomy", width=34),
    Col("duration", "เวลาผ่าตัดที่จอง (นาที)", "int", required=True, keep_blank=True, example="90", width=14),
    Col("stddev", "คลาดเคลื่อน ± (นาที)", "int", keep_blank=True, example="15", width=12),
    Col("role_ANESTHESIOLOGIST", "วิสัญญีแพทย์ (คน)", "int", keep_blank=True, example="1", width=12),
    Col("role_SCRUB_NURSE", "พยาบาลส่งเครื่องมือ (คน)", "int", keep_blank=True, example="1", width=14),
    Col("role_CIRCULATING_NURSE", "พยาบาลช่วยรอบนอก (คน)", "int", keep_blank=True, example="1", width=14),
    Col("role_SURGEON", "ศัลยแพทย์ผู้ช่วย (คน)", "int", keep_blank=True, example="0", width=12),
    Col("equipment", "เครื่องมือที่ใช้ทุกครั้ง", "counts",
        help="ชื่อชนิดเครื่องมือ x จำนวน · คั่นด้วย ; เช่น LAPAROSCOPY_TRAY x2; C_ARM", width=34),
    Col("notes", "หมายเหตุ", help="เช่น ท่าผู้ป่วย ไหม ขนาดถุงมือ", width=30),
), "หนึ่งแถวต่อ แพทย์ 1 ท่าน × หัตถการ 1 อย่าง · ทีมไม่นับศัลยแพทย์เจ้าของเคส (0 = ไม่ต้องใช้)")

SHEETS = (ROOMS, BEDS, TYPES, UNITS, STAFF, SUPPLIES, PROCEDURES, CARDS)  # import order: references first
KEYS = {  # columns a sheet must have: rows are matched to stored records by these
    "rooms": ("code",), "beds": ("code",), "types": ("name",), "units": ("serial",), "staff": ("code",),
    "supplies": ("code",), "procedures": ("name",), "cards": ("surgeon_code", "procedure"),
}
NEW_DEFAULTS = {
    "rooms": {"type": "STANDARD", "capabilities": [], "base_turnover_mins": 30, "infectious_turnover_mins": 60,
              "status": "AVAILABLE", "elective_start": None, "elective_end": None},
    "beds": {"type": "PACU", "status": "AVAILABLE", "note": None},
    "types": {"category": "INSTRUMENT_TRAY", "requires_sterilization": True, "sterilization_mins": 120},
    "units": {"location": None},
    "staff": {"specialty": None, "skills": [], "max_hours_per_week": 48.0, "max_hours_per_day": 10.0,
              "min_rest_hours": 8.0, "active": True},
    "supplies": {"category": "CONSUMABLE", "unit": "ชิ้น", "on_hand": 0, "reorder_level": 0, "active": True},
    "procedures": {"code": None, "duration": 60, "infection": "CLEAN", "icu": False, "blood": False,
                   "recovery_mins": None, "caps": [], "skills": {}, "supplies": {}},
    "cards": {"stddev": 15, "role_ANESTHESIOLOGIST": 1, "role_SCRUB_NURSE": 1, "role_CIRCULATING_NURSE": 1,
              "role_SURGEON": 0, "equipment": {}, "notes": None},
}


class _CellError(Exception):
    pass


# ---------------------------------------------------------------- values

def _squash(text: str) -> str:
    return re.sub(r"\s+", "", str(text)).lower()


def _blank(value) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _text(value) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M")
    return str(value).strip()


def _split(text: str, seps: str = r"[;\n]") -> list[str]:
    return [p.strip() for p in re.split(seps, text) if p.strip()]


def parse_cell(col: Col, value):
    """A non-blank cell -> the value the save code takes. Raises _CellError."""
    if col.kind in ("text", "info"):
        return _text(value)
    if col.kind in ("int", "float"):
        try:
            number = float(str(value).replace(",", "").strip()) if not isinstance(value, (int, float)) else float(value)
        except ValueError:
            raise _CellError(f"{col.label} ต้องเป็นตัวเลข (ได้ {_text(value)!r})") from None
        if col.kind == "int":
            if not number.is_integer():
                raise _CellError(f"{col.label} ต้องเป็นจำนวนเต็ม (ได้ {_text(value)})")
            return int(number)
        return number
    if col.kind == "bool":
        word = _squash(_text(value))
        if word in YES:
            return True
        if word in NO:
            return False
        raise _CellError(f"{col.label} ใส่ ใช่ หรือ ไม่ (ได้ {_text(value)!r})")
    if col.kind == "enum":
        word = _text(value)
        if word.upper() in col.choices:
            return word.upper()
        for alias, code in col.aliases:
            if _squash(alias) == _squash(word):
                return code
        raise _CellError(f"{col.label} ไม่ถูกต้อง: {word!r} (ใช้ได้: {', '.join(col.choices)})")
    if col.kind == "list":
        return sorted({p.strip().upper() for p in re.split(r"[,;\n]", _text(value)) if p.strip()})
    if col.kind == "time":
        return _clock(col, value)
    if col.kind == "skills":
        out: dict[str, list[str]] = {}
        for part in _split(_text(value)):
            if ":" not in part:
                raise _CellError(f"{col.label}: ใช้รูปแบบ บทบาท: ทักษะ (ได้ {part!r})")
            role, names = part.split(":", 1)
            role = role.strip().upper()
            tags = sorted({t.strip().upper() for t in names.split(",") if t.strip()})
            if tags:
                out[role] = sorted(set(out.get(role, [])) | set(tags))
        return out
    if col.kind == "counts":
        out_counts: dict[str, int] = {}
        for part in _split(_text(value)):
            m = re.match(r"^(.*?)\s*[x×*]\s*(\d+)$", part, re.IGNORECASE)
            name, count = (m.group(1).strip(), int(m.group(2))) if m else (part, 1)
            if not name:
                raise _CellError(f"{col.label}: ไม่เข้าใจ {part!r}")
            out_counts[name] = out_counts.get(name, 0) + count
        return out_counts
    raise AssertionError(col.kind)


def _clock(col: Col, value) -> str:
    if isinstance(value, datetime):
        value = value.time()
    if isinstance(value, time):
        return value.strftime("%H:%M")
    if isinstance(value, (int, float)) and 0 <= value < 1:  # Excel stores a time as a fraction of a day
        minutes = round(value * 1440)
        return f"{minutes // 60:02d}:{minutes % 60:02d}"
    m = re.match(r"^(\d{1,2})[:.](\d{2})(?::00)?$", _text(value))
    if not m or int(m.group(1)) > 23 or int(m.group(2)) > 59:
        raise _CellError(f"{col.label} ใช้รูปแบบ HH:MM (ได้ {_text(value)!r})")
    return f"{int(m.group(1)):02d}:{m.group(2)}"


def _empty(col: Col):
    return {"list": [], "skills": {}, "counts": {}}.get(col.kind)


def _norm(value):
    """Comparable form: is this cell different from what is stored?"""
    if value in (None, "", [], {}):
        return None
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return round(float(value), 2)
    if isinstance(value, list):
        return tuple(sorted(str(v).upper() for v in value))
    if isinstance(value, dict):
        return tuple(sorted((str(k).lower(), _norm(v)) for k, v in value.items() if _norm(v) not in (None, 0.0)))
    return str(value).strip()


def cell_text(col: Col, value):
    """What the downloaded form shows for a stored value."""
    if value in (None, [], {}):
        return None
    if col.kind == "bool":
        return _yes_no(bool(value))
    if col.kind == "list":
        return ", ".join(value)
    if col.kind == "skills":
        return "; ".join(f"{role}: {', '.join(tags)}" for role, tags in value.items() if tags)
    if col.kind == "counts":
        return "; ".join(f"{name} x{n}" if n != 1 else name for name, n in value.items() if n)
    if col.kind == "float" and float(value).is_integer():
        return int(value)
    return value


def _shown(col: Col, value) -> str:
    text = cell_text(col, value)
    return "(ว่าง)" if text in (None, "") else str(text)


# --------------------------------------------------------- current records

def _ready(admin: HospitalAdmin, table: str, column: str) -> bool:
    with admin.store._transaction(write=False) as conn:
        return (table, column) in admin._columns(conn)


def current(admin: HospitalAdmin, kind: str) -> list[dict]:
    """The stored records of one sheet, in the sheet's own terms (plus ``_id`` where useful)."""
    if kind == "rooms":
        return admin.rooms()["rooms"]
    if kind == "beds":
        return [{k: b[k] for k in ("code", "type", "status", "note")} for b in admin.beds()["beds"]]
    if kind == "types":
        return [{"_id": t["id"], "name": t["name"], "category": t["category"],
                 "requires_sterilization": t["requires_sterilization"], "sterilization_mins": t["sterilization_mins"]}
                for t in admin.equipment()["types"]]
    if kind == "units":
        return [{"serial": u["serial"], "type": u["type"], "location": u["location"]} for u in admin.equipment()["units"]]
    if kind == "staff":
        return [{"_id": p["id"], **{k: p[k] for k in ("code", "name", "role", "specialty", "skills", "max_hours_per_week",
                                                     "max_hours_per_day", "min_rest_hours", "active")}}
                for p in admin.staff()["staff"]]
    if kind == "supplies":
        if not _ready(admin, "supply_items", "item_id"):
            return []
        return [{**{k: i[k] for k in ("code", "name", "category", "unit", "on_hand", "reorder_level", "active")},
                 "on_hand_now": i["on_hand"]} for i in admin.inventory()["items"]]
    if kind == "procedures":
        supplies = {}
        if _ready(admin, "supply_items", "item_id"):
            supplies = {p["id"]: p["supplies"] for p in admin.inventory()["procedures"]}
        return [{"_id": p["id"], "code": p["code"], "name": p["name"], "duration": p["duration"],
                 "infection": p["infection"], "icu": p["icu"], "blood": p.get("blood", False),
                 "recovery_mins": p["recovery_mins"], "caps": p["caps"],
                 "skills": p["skills"], "supplies": supplies.get(p["id"], {})}
                for p in admin.procedures()["procedures"]]
    if kind == "cards":
        codes = {p["id"]: p["code"] for p in admin.staff()["staff"]}
        out = []
        for c in admin.cards()["cards"]:
            equipment: dict[str, int] = {}
            for e in c["equipment"]:
                equipment[e["name"]] = equipment.get(e["name"], 0) + e["count"]
            out.append({"_id": c["id"], "_surgeon_id": c["surgeon_id"], "_procedure_id": c["procedure_id"],
                        "surgeon_code": codes.get(c["surgeon_id"], ""), "surgeon": c["surgeon"],
                        "procedure": c["procedure"], "duration": c["duration"], "stddev": c["stddev"],
                        **{f"role_{r}": int((c["roles"] or {}).get(r, 0)) for r in CARD_ROLES},
                        "equipment": equipment, "notes": c["notes"]})
        return out
    raise AssertionError(kind)


# -------------------------------------------------------------- download

def build_workbook(admin: HospitalAdmin, *, blank: bool = False) -> bytes:
    """The hospital data as an Excel file (or the empty form), with a how-to sheet."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    head_font, head_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="0E7490")
    wb = Workbook()
    guide = wb.active
    guide.title = GUIDE
    guide.append(["ข้อมูลโรงพยาบาลสำหรับ ORchestraAI"])
    guide["A1"].font = Font(bold=True, size=14)
    for line in [
        "1. แต่ละแผ่นคือข้อมูลหนึ่งประเภท แถวแรกคือหัวคอลัมน์ — อย่าแก้ชื่อหัวคอลัมน์ (ลบคอลัมน์ที่ไม่ใช้ได้ ระบบคงค่าเดิมไว้)",
        "2. คอลัมน์ที่มี * ต้องกรอกเมื่อเพิ่มรายการใหม่ · ช่องที่เป็นตัวเลือกมีรายการให้เลือก (พิมพ์ภาษาไทยได้ เช่น ศัลยแพทย์)",
        "3. ระบบจับคู่กับข้อมูลเดิมด้วยรหัส (หรือชื่อ): แถวที่ตรงกันจะถูกแก้ไข แถวใหม่จะถูกเพิ่ม รายการที่ไม่มีในไฟล์จะไม่ถูกลบ",
        "4. ช่องว่าง: ข้อความ รายการ และเวลา = ล้างค่า · ตัวเลขและตัวเลือก = คงค่าเดิม (รายการใหม่ใช้ค่ามาตรฐาน)",
        "5. อัปโหลดที่หน้า ข้อมูลโรงพยาบาล → นำเข้าข้อมูล ระบบจะตรวจทุกแถวและแสดงผลก่อน ยังไม่บันทึกจนกว่าจะกดยืนยัน",
        "6. ลำดับที่ระบบนำเข้า: ห้อง → เตียง → ชนิดเครื่องมือ → เครื่องรายชิ้น → บุคลากร → วัสดุ → หัตถการ → preference card",
        "7. บันทึกเป็น .xlsx (หรือส่งทีละแผ่นเป็น CSV UTF-8)",
    ]:
        guide.append([line])
    guide.append([])
    guide.append(["แผ่น", "คอลัมน์", "ต้องกรอก", "ความหมาย / ค่าที่ใช้ได้", "ตัวอย่าง"])
    for cell in guide[guide.max_row]:
        cell.font, cell.fill = head_font, head_fill
    for sheet in SHEETS:
        guide.append([sheet.title, "", "", sheet.help, ""])
        guide.cell(guide.max_row, 1).font = Font(bold=True)
        for col in sheet.cols:
            guide.append(["", col.label, "*" if col.required else "", col.help, col.example])
    for letter, width in zip("ABCDE", (18, 34, 9, 80, 30)):
        guide.column_dimensions[letter].width = width
    for (cell,) in guide.iter_rows(min_row=11, min_col=4, max_col=4):
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    for sheet in SHEETS:
        ws = wb.create_sheet(sheet.title)
        ws.append([col.label + (" *" if col.required else "") for col in sheet.cols])
        for cell in ws[1]:
            cell.font, cell.fill = head_font, head_fill
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[1].height = 34
        ws.freeze_panes = "A2"
        rows = [] if blank else current(admin, sheet.kind)
        for record in rows:
            ws.append([cell_text(col, record.get(col.key)) if col.download else None for col in sheet.cols])
        last = max(FORM_ROWS, len(rows) + 200)
        for i, col in enumerate(sheet.cols, start=1):
            letter = get_column_letter(i)
            ws.column_dimensions[letter].width = col.width
            if col.kind not in ("int", "float"):  # keep codes like 06.4 and times like 08:00 as typed
                ws.column_dimensions[letter].number_format = "@"
                for (cell,) in ws.iter_rows(min_row=2, max_row=len(rows) + 1, min_col=i, max_col=i):
                    cell.number_format = "@"
            choices = list(col.choices) if col.kind == "enum" else ["ใช่", "ไม่"] if col.kind == "bool" else None
            if choices:
                dv = DataValidation(type="list", formula1='"' + ",".join(choices) + '"', allow_blank=True)
                dv.showErrorMessage = False  # Thai words are accepted too; the import checks them
                ws.add_data_validation(dv)
                dv.add(f"{letter}2:{letter}{last}")
            if col.kind == "info":
                for (cell,) in ws.iter_rows(min_row=2, max_row=len(rows) + 1, min_col=i, max_col=i):
                    cell.font = Font(color="64748B")
    for ws in wb.worksheets:  # print across one page width
        ws.page_setup.orientation = "landscape"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    out = io.BytesIO()
    as_text(wb)
    wb.save(out)
    return out.getvalue()


# ---------------------------------------------------------------- upload

@dataclass
class Table:
    sheet: Sheet
    name: str  # sheet or file name, for messages
    cols: list  # Col per header position (None = ignored column)
    rows: list  # (excel row number, [cell values])
    unknown: list  # header texts that matched nothing


def _header_key(text) -> str:
    return _squash(str(text or "").replace("*", ""))


def _match(sheet: Sheet, header: list) -> tuple[list, list, int]:
    by_name = {}
    for col in sheet.cols:
        by_name[_header_key(col.label)] = col
        by_name[_header_key(col.key)] = col
    cols, unknown, hits = [], [], 0
    for text in header:
        col = by_name.get(_header_key(text)) if not _blank(text) else None
        if col is not None and col not in cols:
            cols.append(col)
            hits += 1
        else:
            cols.append(None)
            if not _blank(text):
                unknown.append(_text(text))
    return cols, unknown, hits


def _table(name: str, raw_rows: list[list]) -> Table | None:
    """Which sheet a grid of cells is (by its name, else by its header row)."""
    wanted = next((s for s in SHEETS if _squash(s.title) == _squash(name)), None)
    best = None
    for index, header in enumerate(raw_rows[:10]):  # the header may sit under a title row
        for sheet in ([wanted] if wanted else SHEETS):
            cols, unknown, hits = _match(sheet, list(header))
            has_keys = all(any(c is not None and c.key == k for c in cols) for k in KEYS[sheet.kind])
            if hits >= 2 and (wanted or has_keys) and (best is None or hits > best[0]):
                best = (hits, index, sheet, cols, unknown)
    if best is None:
        return None
    _, index, sheet, cols, unknown = best
    rows = [(index + 2 + n, list(r)) for n, r in enumerate(raw_rows[index + 1:]) if any(not _blank(v) for v in r)]
    return Table(sheet, name, cols, rows, unknown)


def read_tables(filename: str, content: bytes) -> tuple[list[Table], list[str]]:
    """Parse an uploaded .xlsx (every sheet) or .csv (one sheet)."""
    if not content:
        raise StoreError("ไฟล์ว่างเปล่า")
    if len(content) > MAX_BYTES:
        raise StoreError(f"ไฟล์ใหญ่เกิน {MAX_BYTES // (1024 * 1024)} MB")
    name = (filename or "").lower()
    grids: list[tuple[str, list[list]]] = []
    if content[:2] == b"PK" or name.endswith(".xlsx"):
        grids = _xlsx_grids(content)
    elif name.endswith(".xls"):
        raise StoreError("ไฟล์ .xls (Excel รุ่นเก่า) ยังไม่รองรับ กรุณาบันทึกเป็น .xlsx หรือ CSV UTF-8")
    else:
        grids = [(filename or "CSV", _csv_grid(content))]
    tables, warnings = [], []
    seen: dict[str, str] = {}
    for sheet_name, grid in grids:
        if _squash(sheet_name) == _squash(GUIDE):
            continue
        if len(grid) > MAX_ROWS + 10:
            raise StoreError(f"แผ่น {sheet_name} มีเกิน {MAX_ROWS} แถว กรุณาแบ่งเป็นหลายไฟล์")
        table = _table(sheet_name, grid)
        if table is None:
            if any(any(not _blank(v) for v in r) for r in grid):
                warnings.append(f"ข้ามแผ่น {sheet_name}: ไม่พบหัวคอลัมน์ที่รู้จัก")
            continue
        if table.sheet.kind in seen:
            raise StoreError(f"มีข้อมูล{table.sheet.title}สองแผ่น ({seen[table.sheet.kind]}, {sheet_name}) กรุณารวมเป็นแผ่นเดียว")
        seen[table.sheet.kind] = sheet_name
        tables.append(table)
    if not tables:
        raise StoreError("ไม่พบข้อมูลที่นำเข้าได้ในไฟล์นี้ (ใช้ไฟล์ที่ดาวน์โหลดจากหน้านี้เป็นแบบ)")
    order = [s.kind for s in SHEETS]
    return sorted(tables, key=lambda t: order.index(t.sheet.kind)), warnings


def _xlsx_grids(content: bytes) -> list[tuple[str, list[list]]]:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as z:
            infos = z.infolist()
            if len(infos) > 1000 or sum(i.file_size for i in infos) > MAX_UNZIPPED:
                raise StoreError("ไฟล์ Excel ใหญ่เกินไป")
    except zipfile.BadZipFile as exc:
        raise StoreError("เปิดไฟล์ Excel ไม่ได้ (ไฟล์เสียหรือไม่ใช่ .xlsx)") from exc
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        if len(wb.worksheets) > MAX_SHEETS:
            raise StoreError(f"ไฟล์มีแผ่นงานเกิน {MAX_SHEETS} แผ่น")
        grids = []
        for ws in wb.worksheets:
            rows = []
            # max_col: a sheet that claims thousands of columns must not be read (and padded) in full
            for row in ws.iter_rows(max_col=60, values_only=True):
                rows.append(list(row))
                if len(rows) > MAX_ROWS + 10:
                    break
            grids.append((ws.title, rows))
        wb.close()
    except StoreError:
        raise
    except Exception as exc:  # openpyxl raises many kinds of errors for damaged files
        raise StoreError(f"เปิดไฟล์ Excel ไม่ได้: {exc}") from exc
    return grids


def _csv_grid(content: bytes) -> list[list]:
    for encoding in ("utf-8-sig", "cp874"):  # CSV UTF-8, or Thai Windows Excel's default
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise StoreError("อ่านไฟล์ CSV ไม่ได้ กรุณาบันทึกเป็น CSV UTF-8")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    try:
        return [row[:60] for row in itertools.islice(csv.reader(io.StringIO(text), dialect), MAX_ROWS + 11)]
    except csv.Error as exc:  # e.g. a field longer than csv.field_size_limit()
        raise StoreError(f"อ่านไฟล์ CSV ไม่ได้: {exc}") from exc


# ---------------------------------------------------------------- import

class _Undo(Exception):
    """Leave the import transaction without saving (preview, or errors)."""


class _Run:
    """One sheet's rows matched to stored records and saved through the admin service."""

    def __init__(self, admin: HospitalAdmin, table: Table, actor: dict | None, filename: str) -> None:
        self.admin, self.table, self.actor, self.filename = admin, table, actor, filename
        self.sheet = table.sheet
        self.present = [c for c in table.cols if c is not None]
        self.result = {"kind": self.sheet.kind, "title": self.sheet.title, "source": table.name,
                       "new": 0, "changed": 0, "unchanged": 0, "errors": [], "rows": [], "warnings": []}
        self.records = current(admin, self.sheet.kind)
        self.lookups()

    # --- matching
    def lookups(self) -> None:
        kind = self.sheet.kind
        lower = lambda s: str(s or "").strip().lower()  # noqa: E731
        if kind == "procedures":
            self.by_code = {lower(r["code"]): r for r in self.records if r["code"]}
            self.by_name = {lower(r["name"]): r for r in self.records}
            self.supply_codes = {lower(i["code"]): i["code"] for i in current(self.admin, "supplies")}
        elif kind == "cards":
            staff = current(self.admin, "staff")
            self.surgeons = {lower(p["code"]): p for p in staff if p["role"] == "SURGEON"}
            self.others = {lower(p["code"]): p for p in staff if p["role"] != "SURGEON"}
            procs = current(self.admin, "procedures")
            self.proc_code = {lower(p["code"]): p for p in procs if p["code"]}
            self.proc_name = {lower(p["name"]): p for p in procs}
            self.type_ids = {lower(t["name"]): t for t in current(self.admin, "types")}
            self.by_pair = {(r["_surgeon_id"], r["_procedure_id"]): r for r in self.records}
        else:
            if kind == "units":
                self.type_ids = {lower(t["name"]): t for t in current(self.admin, "types")}
            key = {"types": "name", "units": "serial"}.get(kind, "code")
            self.by_key = {lower(r[key]): r for r in self.records}

    def find(self, cells: dict):
        """(stored record or None, key text shown in the report)."""
        kind, lower = self.sheet.kind, (lambda s: str(s or "").strip().lower())
        if kind == "procedures":
            code, name = cells.get("code"), cells.get("name")
            hit = (self.by_code.get(lower(code)) if not _blank(code) else None) or \
                (self.by_name.get(lower(name)) if not _blank(name) else None)
            return hit, _text(name if not _blank(name) else code or "")
        if kind == "cards":
            code, proc = cells.get("surgeon_code"), cells.get("procedure")
            if _blank(code) or _blank(proc):
                raise StoreError("ต้องกรอกรหัสศัลยแพทย์และหัตถการ")
            surgeon = self.surgeons.get(lower(code))
            if surgeon is None:
                other = self.others.get(lower(code))
                raise StoreError(f"{_text(code)} ไม่ใช่ศัลยแพทย์ ({other['role']})" if other
                                 else f"ไม่พบศัลยแพทย์รหัส {_text(code)} (เพิ่มในแผ่นบุคลากรก่อน)")
            procedure = self.proc_code.get(lower(proc)) or self.proc_name.get(lower(proc))
            if procedure is None:
                raise StoreError(f"ไม่พบหัตถการ {_text(proc)} (เพิ่มในแผ่นหัตถการก่อน)")
            self.pair = (surgeon["_id"], procedure["_id"])
            self.names = (surgeon["code"], procedure["name"])  # as stored, however the row wrote them
            return self.by_pair.get(self.pair), f"{surgeon['name']} · {procedure['name']}"
        key = {"types": "name", "units": "serial"}.get(kind, "code")
        if _blank(cells.get(key)):
            col = next(c for c in self.sheet.cols if c.key == key)
            raise StoreError(f"ต้องกรอก{col.label}")
        return self.by_key.get(lower(cells[key])), _text(cells[key])

    # --- one row
    def row(self, number: int, raw: list, seen: dict) -> None:
        cells = {col.key: raw[i] if i < len(raw) else None for i, col in enumerate(self.table.cols) if col is not None}
        key_text = " · ".join(_text(cells[k]) for k in KEYS[self.sheet.kind] if not _blank(cells.get(k)))
        try:
            stored, key_text = self.find(cells)
            ident = ("pair", self.pair) if self.sheet.kind == "cards" else key_text.lower()
            if ident in seen:
                raise StoreError(f"ซ้ำกับแถว {seen[ident]}")
            seen[ident] = number
            record = dict(stored) if stored else dict(NEW_DEFAULTS[self.sheet.kind])
            for col in self.present:
                if col.kind == "info":
                    continue
                value = cells.get(col.key)
                if _blank(value):
                    if stored is None and col.required:
                        raise StoreError(f"ต้องกรอก{col.label}")
                    if not col.keep_blank:
                        record[col.key] = _empty(col)
                    continue
                try:
                    record[col.key] = parse_cell(col, value)
                except _CellError as exc:
                    raise StoreError(str(exc)) from None
            for k in KEYS[self.sheet.kind]:  # matched ignoring case: keep the stored spelling of the key
                if stored and str(record.get(k) or "").strip().lower() == str(stored.get(k) or "").strip().lower():
                    record[k] = stored[k]
            if stored is None:
                missing = [c.label for c in self.sheet.cols if c.required and c.key not in record]
                if missing:
                    raise StoreError(f"ต้องมีคอลัมน์ {', '.join(missing)} สำหรับรายการใหม่")
            self.resolve(record)
            changes = [] if stored is None else [
                f"{c.label}: {_shown(c, stored.get(c.key))} → {_shown(c, record.get(c.key))}"
                for c in self.present if c.kind != "info" and _norm(stored.get(c.key)) != _norm(record.get(c.key))]
            if stored is not None and not changes:
                self.result["unchanged"] += 1
                return
            with self.admin.store._transaction(write=True):  # this row on its own: a failure undoes only it
                self.save(record, stored)
            action = "new" if stored is None else "changed"
            self.result[action] += 1
            self.result["rows"].append({"row": number, "key": key_text, "action": action, "changes": changes[:8]})
        except (StoreError, psycopg.Error, ValueError, KeyError) as exc:
            message = str(exc).splitlines()[0] if str(exc) else type(exc).__name__
            self.result["errors"].append({"row": number, "key": key_text, "message": message})

    def resolve(self, record: dict) -> None:
        """Names in the sheet -> the codes and ids the save code takes."""
        lower = lambda s: str(s or "").strip().lower()  # noqa: E731
        if self.sheet.kind == "units":
            t = self.type_ids.get(lower(record.get("type")))
            if t is None:
                raise StoreError(f"ไม่พบชนิดเครื่องมือ {record.get('type')} (เพิ่มในแผ่นชนิดเครื่องมือก่อน)")
            record["type"] = t["name"]
            record["_type_id"] = t["_id"]
        elif self.sheet.kind == "procedures":
            record["supplies"] = {self.supply_codes.get(lower(k), k): v for k, v in (record.get("supplies") or {}).items()}
        elif self.sheet.kind == "cards":
            record["surgeon_code"], record["procedure"] = self.names
            equipment = {}
            for name, count in (record.get("equipment") or {}).items():
                t = self.type_ids.get(lower(name))
                if t is None:
                    parts = [p.strip() for p in name.split(",") if p.strip()]  # "A, B" typed with commas
                    if len(parts) > 1 and all(lower(p) in self.type_ids for p in parts):
                        for p in parts:
                            real = self.type_ids[lower(p)]["name"]
                            equipment[real] = equipment.get(real, 0) + 1
                        continue
                    raise StoreError(f"ไม่พบชนิดเครื่องมือ {name}")
                equipment[t["name"]] = equipment.get(t["name"], 0) + count
            record["equipment"] = equipment

    def save(self, r: dict, stored: dict | None) -> None:
        a, kind, actor, new = self.admin, self.sheet.kind, self.actor, stored is None
        if kind == "rooms":
            a.save_room(r, code=None if new else stored["code"], actor=actor)
        elif kind == "beds":
            a.save_bed(r, code=None if new else stored["code"], actor=actor)
        elif kind == "types":
            a.save_type(r, type_id=None if new else stored["_id"], actor=actor)
        elif kind == "units":
            a.save_unit({"serial": r["serial"], "type_id": r["_type_id"], "location": r.get("location")},
                        serial=None if new else stored["serial"], actor=actor)
        elif kind == "staff":
            a.save_staff(r, staff_id=None if new else stored["_id"], actor=actor)
        elif kind == "supplies":
            if new:
                a.save_item(r, actor=actor)
                return
            if any(_norm(stored.get(k)) != _norm(r.get(k)) for k in ("name", "category", "unit", "reorder_level", "active")):
                a.save_item(r, code=stored["code"], actor=actor)
            if _norm(stored.get("on_hand")) != _norm(r.get("on_hand")):
                a.count_stock(stored["code"], {"on_hand": r["on_hand"], "note": f"นำเข้าจาก {self.filename}"}, actor=actor)
        elif kind == "procedures":
            pid = a.save_procedure(r, procedure_id=None if new else stored["_id"], actor=actor)["id"]
            old = {} if new else stored.get("supplies") or {}
            if _norm(old) != _norm(r.get("supplies")):
                if not hasattr(a, "set_procedure_supplies"):
                    raise StoreError("ยังตั้งวัสดุของหัตถการไม่ได้")
                a.set_procedure_supplies(pid, {"supplies": r.get("supplies") or {}}, actor=actor)
        elif kind == "cards":
            types = {t["name"]: t["_id"] for t in self.type_ids.values()}
            a.save_card({"surgeon_id": self.pair[0], "procedure_id": self.pair[1], "duration": r["duration"],
                         "stddev": r["stddev"], "roles": {role: r[f"role_{role}"] for role in CARD_ROLES},
                         "equipment": [{"type_id": types[n], "count": c} for n, c in r["equipment"].items()],
                         "notes": r.get("notes")}, card_id=None if new else stored["_id"], actor=actor)

    def run(self) -> dict:
        sheet = self.sheet
        if sheet.needs and not _ready(self.admin, *sheet.needs):
            self.result["warnings"].append(f"ข้ามแผ่น{sheet.title}: ฐานข้อมูลยังไม่ได้รัน migration ของคลังวัสดุ")
            return self.result
        if sheet.kind == "procedures" and any(c.key == "supplies" for c in self.present) \
                and not _ready(self.admin, "supply_items", "item_id"):
            self.present = [c for c in self.present if c.key != "supplies"]
            self.table.cols = [None if c is not None and c.key == "supplies" else c for c in self.table.cols]
            self.result["warnings"].append("ข้ามคอลัมน์วัสดุที่ใช้: ฐานข้อมูลยังไม่ได้รัน migration ของคลังวัสดุ")
        missing = [c.label for c in sheet.cols if c.key in KEYS[sheet.kind] and c not in self.present]
        if missing:
            self.result["errors"].append({"row": 1, "key": "", "message": f"ไม่พบคอลัมน์ {', '.join(missing)}"})
            return self.result
        if self.table.unknown:
            self.result["warnings"].append(f"ไม่รู้จักคอลัมน์ (ข้าม): {', '.join(self.table.unknown[:8])}")
        seen: dict = {}
        for number, raw in self.table.rows:
            self.row(number, raw, seen)
        return self.result


def import_file(admin: HospitalAdmin, filename: str, content: bytes, *, commit: bool,
                actor: dict | None = None) -> dict:
    """Check (and, when ``commit`` and nothing is wrong, save) an uploaded workbook."""
    tables, warnings = read_tables(filename, content)
    sheets: list[dict] = []
    affected: list[dict] = []
    committed = False
    try:
        with admin.store.batch() as conn:
            admin._batch, admin._schema = [], None
            for table in tables:
                sheets.append(_Run(admin, table, actor, filename).run())
            errors = sum(len(s["errors"]) for s in sheets)
            saved = sum(s["new"] + s["changed"] for s in sheets)
            if errors or not commit or not saved:
                raise _Undo
            admin._batch = None
            lines = [f"{s['title']}: เพิ่ม {s['new']} แก้ไข {s['changed']} ไม่เปลี่ยน {s['unchanged']}" for s in sheets]
            admin._log(conn, f"นำเข้าข้อมูลจาก {filename[:80]}: เพิ่ม {sum(s['new'] for s in sheets)} "
                             f"แก้ไข {sum(s['changed'] for s in sheets)} รายการ", lines, actor)
            today = admin.db.today()
            affected = admin._affected(conn, today, today + timedelta(days=CHECK_DAYS))
            committed = True
    except _Undo:
        pass
    finally:
        admin._batch, admin._schema = None, None
    return {
        "filename": filename, "committed": committed, "sheets": sheets, "warnings": warnings, "affected": affected,
        "new": sum(s["new"] for s in sheets), "changed": sum(s["changed"] for s in sheets),
        "unchanged": sum(s["unchanged"] for s in sheets), "errors": sum(len(s["errors"]) for s in sheets),
    }


# --------------------------------------------------- safe downloads (no formulas)

FORMULA_START = ("=", "+", "-", "@", "\t", "\r")


def as_text(wb) -> None:
    """Keep typed text as text: openpyxl stores any string starting with "=" as a formula, which would run
    when someone opens the download (a name or note like =HYPERLINK(...) must stay plain text)."""
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    cell.data_type = "s"


def csv_safe(value):
    """A CSV cell that Excel will not read as a formula (numbers stay numbers)."""
    if isinstance(value, str) and value.startswith(FORMULA_START):
        try:
            float(value)
            return value
        except ValueError:
            return "'" + value
    return value


def list_formula(choices) -> str:
    """A drop-down list for data validation: quotes doubled, codes holding a comma left out."""
    return '"' + ",".join(str(c).replace('"', '""') for c in choices if "," not in str(c))[:250] + '"'


# --------------------------------------------------- demo data, start over

def _demo_where() -> tuple[str, tuple]:
    return ("(r.patient_hn LIKE 'HIST-%%' AND r.raw_nlp_input = %s) OR r.patient_name LIKE %s", (HISTORY_TAG, DEMO_NAME))


WIPE_ORDER = (  # children before parents
    "resource_allocations", "supply_reservations", "scheduled_surgeries", "surgery_requests",
    "surgeon_preference_cards", "procedure_supplies", "supply_deliveries", "supply_items", "procedures",
    "equipment_maintenance", "equipment_units", "equipment_types", "room_closures", "operating_rooms",
    "recovery_beds", "team_synergy", "staff_schedules", "staff",
)
COUNTED = {"operating_rooms": "ห้องผ่าตัด", "recovery_beds": "เตียงพักฟื้น", "staff": "บุคลากร",
           "equipment_units": "เครื่องมือ", "procedures": "หัตถการ", "surgeon_preference_cards": "preference card",
           "supply_items": "วัสดุ", "surgery_requests": "คำขอผ่าตัด", "scheduled_surgeries": "เคสในตาราง"}


def _tables(conn) -> set[str]:
    return {r["table_name"] for r in conn.execute(
        "SELECT table_name FROM information_schema.tables WHERE table_schema = 'public'")}


def data_status(admin: HospitalAdmin) -> dict:
    """How much demo data and data in total there is (for the clean-up buttons)."""
    where, args = _demo_where()
    with admin.store._transaction(write=False) as conn:
        tables = _tables(conn)
        row = conn.execute(
            f"""SELECT count(*) FILTER (WHERE r.patient_hn LIKE 'HIST-%%') AS history,
                       count(*) FILTER (WHERE r.patient_name LIKE %s) AS cases
                FROM surgery_requests r WHERE {where}""", (DEMO_NAME, *args)).fetchone()
        counts = {label: conn.execute(f"SELECT count(*) AS n FROM {table}").fetchone()["n"]  # noqa: S608 (fixed names)
                  for table, label in COUNTED.items() if table in tables}
        linked = conn.execute("SELECT count(*) AS n FROM app_users WHERE staff_id IS NOT NULL").fetchone()["n"] \
            if "app_users" in tables else 0
    return {"demo_history": row["history"], "demo_cases": row["cases"], "counts": counts, "linked_users": linked,
            "wipe_phrase": WIPE_PHRASE}


def delete_demo_cases(admin: HospitalAdmin, *, actor: dict | None = None) -> dict:
    """Remove the demo day's cases and the invented history (HN HIST-...); real cases stay."""
    where, args = _demo_where()
    with admin.store._transaction(write=True) as conn:
        ids = [r["request_id"] for r in conn.execute(f"SELECT r.request_id FROM surgery_requests r WHERE {where}", args)]
        booked = conn.execute("DELETE FROM scheduled_surgeries WHERE request_id = ANY(%s)", (ids,)).rowcount if ids else 0
        removed = conn.execute("DELETE FROM surgery_requests WHERE request_id = ANY(%s)", (ids,)).rowcount if ids else 0
        admin._log(conn, f"ลบเคสตัวอย่างและประวัติตัวอย่าง {removed} รายการ",
                   [f"คำขอ {removed}, เคสในตาราง {booked}"], actor)
    return {"requests": removed, "bookings": booked}


def wipe_hospital(admin: HospitalAdmin, phrase: str, *, actor: dict | None = None) -> dict:
    """Delete every hospital record and booking to start over; user accounts and settings stay."""
    if (phrase or "").strip() != WIPE_PHRASE:
        raise StoreError(f'พิมพ์คำว่า "{WIPE_PHRASE}" เพื่อยืนยัน')
    with admin.store._transaction(write=True) as conn:
        tables = _tables(conn)
        before = {t: conn.execute(f"SELECT count(*) AS n FROM {t}").fetchone()["n"]  # noqa: S608 (fixed names)
                  for t in COUNTED if t in tables}
        unlinked = conn.execute("SELECT count(*) AS n FROM app_users WHERE staff_id IS NOT NULL").fetchone()["n"] \
            if "app_users" in tables else 0
        for table in WIPE_ORDER:
            if table in tables:
                conn.execute(f"DELETE FROM {table}")  # noqa: S608 (fixed names)
        lines = [f"{COUNTED[t]} {n}" for t, n in before.items() if n]
        admin._log(conn, "ล้างข้อมูลโรงพยาบาลทั้งหมด (เก็บบัญชีผู้ใช้ไว้)",
                   lines + [f"ยกเลิกการผูกบัญชีผู้ใช้กับรายชื่อบุคลากร {unlinked} บัญชี"], actor)
    return {"deleted": {COUNTED[t]: n for t, n in before.items()}, "unlinked_users": unlinked}
