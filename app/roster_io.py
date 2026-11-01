"""Monthly rosters (shifts, on-call and leave) in and out of Excel.

A roster sheet is people down the side and days across the top; each cell holds
shift codes such as ``ช`` (08:00-16:00), ``ช/บ`` (two shifts), ``OC`` (on-call),
``ล`` (leave) or a time range such as ``07:30-15:30``. The codes are defined in the
workbook's "รหัสเวร" sheet (the hospital's own codes; the last imported set is kept
and used for the next download). Headers may be real dates, ISO / d/m/yyyy text
(Buddhist years too) or day numbers with the month in the sheet name or title.

For every person and day in the file, the cell is the whole roster for that day:
the shifts that start that day are replaced by the cell's (a blank cell = no
shift). People and days that are not in the file are left as they are. Upload
first previews what changes; only a file without errors can be imported, in one
transaction, with one activity entry and a check of the bookings it breaks.
"""
from __future__ import annotations

import io
import json
import re
import uuid
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta

from .admin_store import HospitalAdmin
from .hospital_io import _blank, _csv_grid, _squash, _text, _xlsx_grids
from .store import StoreError

CODES_SHEET = "รหัสเวร"
GUIDE = "วิธีกรอก"
WEEKDAY_ROW = "วันในสัปดาห์"
SETTINGS_KEY = "roster_codes"
TYPES = ("REGULAR", "ON_CALL", "LEAVE", "OFF")
TYPE_TH = {"REGULAR": "เวรปกติ", "ON_CALL": "on-call", "LEAVE": "ลา", "OFF": "หยุด (ไม่มีเวร)"}
TYPE_ALIASES = {"เวร": "REGULAR", "เวรปกติ": "REGULAR", "ปกติ": "REGULAR", "oncall": "ON_CALL", "on-call": "ON_CALL",
                "เวรon-call": "ON_CALL", "ลา": "LEAVE", "วันลา": "LEAVE", "หยุด": "OFF", "วันหยุด": "OFF",
                "ไม่มีเวร": "OFF", "off": "OFF"}
DAY_TH = ["จ", "อ", "พ", "พฤ", "ศ", "ส", "อา"]
MONTHS = {  # Thai full and short names, English names -> month number
    **{n: i for i, n in enumerate(["มกราคม", "กุมภาพันธ์", "มีนาคม", "เมษายน", "พฤษภาคม", "มิถุนายน", "กรกฎาคม",
                                   "สิงหาคม", "กันยายน", "ตุลาคม", "พฤศจิกายน", "ธันวาคม"], 1)},
    **{n: i for i, n in enumerate(["ม.ค.", "ก.พ.", "มี.ค.", "เม.ย.", "พ.ค.", "มิ.ย.", "ก.ค.", "ส.ค.", "ก.ย.", "ต.ค.",
                                   "พ.ย.", "ธ.ค."], 1)},
    **{n: i for i, n in enumerate(["january", "february", "march", "april", "may", "june", "july", "august",
                                   "september", "october", "november", "december"], 1)},
    **{n: i for i, n in enumerate(["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov",
                                   "dec"], 1)},
}
MONTH_TH = ["", "มกราคม", "กุมภาพันธ์", "มีนาคม", "เมษายน", "พฤษภาคม", "มิถุนายน", "กรกฎาคม", "สิงหาคม", "กันยายน",
            "ตุลาคม", "พฤศจิกายน", "ธันวาคม"]
MAX_DAYS = 93  # days in one upload
PAST_DAYS, FUTURE_DAYS = 62, 400  # how far back / ahead an upload may change
MAX_PEOPLE = 1000
SEPARATORS = r"[/+,;\n]"


@dataclass(frozen=True)
class Code:
    code: str
    type: str  # REGULAR / ON_CALL / LEAVE / OFF
    start: time | None = None
    end: time | None = None  # end <= start: ends the next day (00:00-00:00 = the whole day)
    label: str = ""

    def as_row(self) -> list:
        clock = (lambda t: None if t is None else t.strftime("%H:%M"))  # noqa: E731
        return [self.code, self.type, clock(self.start), clock(self.end), self.label]


DEFAULT_CODES = (
    Code("ช", "REGULAR", time(8), time(16), "เวรเช้า"),
    Code("บ", "REGULAR", time(16), time(0), "เวรบ่าย (ถึงเที่ยงคืน)"),
    Code("ด", "REGULAR", time(0), time(8), "เวรดึก (เที่ยงคืนถึง 08:00 ของวันนั้น)"),
    Code("OC", "ON_CALL", time(16), time(8), "on-call ถึง 08:00 ของวันถัดไป"),
    Code("ล", "LEAVE", time(0), time(0), "ลา (ทั้งวัน)"),
    Code("ลป", "LEAVE", time(0), time(0), "ลาป่วย"),
    Code("ลพ", "LEAVE", time(0), time(0), "ลาพักร้อน"),
    Code("ปช", "LEAVE", time(0), time(0), "ประชุม/อบรม"),
    Code("-", "OFF", label="วันหยุด"),
    Code("หยุด", "OFF", label="วันหยุด"),
    Code("OFF", "OFF", label="วันหยุด"),
)


# ---------------------------------------------------------------- codes

def _clock(text, what: str) -> time:
    if isinstance(text, datetime):
        return text.time()
    if isinstance(text, time):
        return text
    if isinstance(text, (int, float)) and 0 <= text < 1:  # Excel time = fraction of a day
        minutes = round(text * 1440) % 1440
        return time(minutes // 60, minutes % 60)
    m = re.match(r"^(\d{1,2})[:.](\d{2})(?::00)?$", _text(text))
    if not m or int(m.group(1)) > 24 or int(m.group(2)) > 59 or (int(m.group(1)) == 24 and int(m.group(2))):
        raise StoreError(f"{what} ใช้รูปแบบ HH:MM (ได้ {_text(text)!r})")
    return time(int(m.group(1)) % 24, int(m.group(2)))


def load_codes(admin: HospitalAdmin) -> list[Code]:
    """The codes of the last imported roster (or the defaults)."""
    with admin.store._transaction(write=False) as conn:
        if ("app_settings", "key") not in admin._columns(conn):
            return list(DEFAULT_CODES)
        row = conn.execute("SELECT value FROM app_settings WHERE key = %s", (SETTINGS_KEY,)).fetchone()
    if row is None:
        return list(DEFAULT_CODES)
    try:
        return [Code(c["code"], c["type"], time.fromisoformat(c["start"]) if c.get("start") else None,
                     time.fromisoformat(c["end"]) if c.get("end") else None, c.get("label") or "")
                for c in json.loads(row["value"])]
    except (ValueError, KeyError, TypeError):
        return list(DEFAULT_CODES)


def _save_codes(admin: HospitalAdmin, conn, codes: list[Code]) -> None:
    if ("app_settings", "key") not in admin._columns(conn):
        return
    value = json.dumps([{"code": c.code, "type": c.type, "start": c.start.isoformat() if c.start else None,
                         "end": c.end.isoformat() if c.end else None, "label": c.label} for c in codes],
                       ensure_ascii=False)
    conn.execute("""INSERT INTO app_settings (key, value) VALUES (%s, %s)
                    ON CONFLICT (key) DO UPDATE SET value = EXCLUDED.value""", (SETTINGS_KEY, value))


def parse_codes(grid: list[list]) -> tuple[list[Code], list[dict]]:
    """The "รหัสเวร" sheet: code, type, start, end, meaning."""
    header_at = next((i for i, row in enumerate(grid[:10]) if row and _squash(_text(row[0] or "")) == "รหัส"), None)
    if header_at is None:
        return [], [{"sheet": CODES_SHEET, "cell": "A1", "message": "ไม่พบหัวคอลัมน์ รหัส / ประเภท / เริ่ม / เลิก"}]
    codes, errors, seen = [], [], {}
    for n, row in enumerate(grid[header_at + 1:], start=header_at + 2):
        cells = (list(row) + [None] * 5)[:5]
        if all(_blank(v) for v in cells):
            continue
        try:
            code = _text(cells[0])
            if not code or len(code) > 12 or re.search(SEPARATORS, code) or re.search(r"\d[:.]\d", code):
                raise StoreError("รหัสเวรต้องมี 1-12 ตัว ห้ามมี / + , ; หรือเวลา")
            kind = _text(cells[1]).upper().replace(" ", "_")
            kind = kind if kind in TYPES else TYPE_ALIASES.get(_squash(_text(cells[1])), "")
            if not kind:
                raise StoreError(f"ประเภทไม่ถูกต้อง: {_text(cells[1])!r} (ใช้ REGULAR, ON_CALL, LEAVE หรือ OFF)")
            if kind == "OFF":
                start = end = None
            else:
                if _blank(cells[2]) or _blank(cells[3]):
                    raise StoreError("ต้องใส่เวลาเริ่มและเวลาเลิก (ลาทั้งวันใส่ 00:00 ถึง 00:00)")
                start, end = _clock(cells[2], "เวลาเริ่ม"), _clock(cells[3], "เวลาเลิก")
            if code.lower() in seen:
                raise StoreError(f"รหัส {code} ซ้ำกับแถว {seen[code.lower()]}")
            seen[code.lower()] = n
            codes.append(Code(code, kind, start, end, _text(cells[4]) if not _blank(cells[4]) else ""))
        except StoreError as exc:
            errors.append({"sheet": CODES_SHEET, "cell": f"A{n}", "message": str(exc)})
    if not codes and not errors:
        errors.append({"sheet": CODES_SHEET, "cell": "A2", "message": "ไม่มีรหัสเวร"})
    return codes, errors


class Codebook:
    """Turns cell text into shifts on a day, and shifts back into cell text."""

    def __init__(self, admin: HospitalAdmin, codes: list[Code]) -> None:
        self.admin = admin
        self.codes = codes
        self.by_code = {c.code.lower(): c for c in codes}
        self.first_of_type = {}
        for c in codes:
            self.first_of_type.setdefault(c.type, c)

    def span(self, day: date, start: time, end: time) -> tuple[datetime, datetime]:
        t0 = self.admin._at(day, start)
        return t0, self.admin._at(day + timedelta(days=1) if end <= start else day, end)

    def shifts(self, day: date, text) -> set[tuple]:
        """Cell -> {(type, start, end, note)}; raises StoreError for an unknown code."""
        out = set()
        for token in [t.strip() for t in re.split(SEPARATORS, _text(text)) if t.strip()]:
            code = self.by_code.get(token.lower())
            if code is not None:
                if code.type != "OFF":  # a leave keeps its kind (ลาป่วย, ประชุม) as the note
                    out.add((code.type, *self.span(day, code.start, code.end),
                             (code.label or None) if code.type == "LEAVE" else None))
                continue
            m = re.match(r"^(?:(\S+?)\s+)?(\d{1,2}[:.]\d{2})\s*[-–]\s*(\d{1,2}[:.]\d{2})$", token)
            if not m:
                raise StoreError(f"ไม่รู้จักรหัสเวร {token!r} (เพิ่มในแผ่น {CODES_SHEET} หรือใช้เวลา เช่น 07:30-15:30)")
            kind = "REGULAR"
            if m.group(1):
                prefix = self.by_code.get(m.group(1).lower())
                if prefix is None or prefix.type == "OFF":
                    raise StoreError(f"ไม่รู้จักรหัสเวร {m.group(1)!r} หน้าเวลา")
                kind = prefix.type
            out.add((kind, *self.span(day, _clock(m.group(2), "เวลาเริ่ม"), _clock(m.group(3), "เวลาเลิก")), None))
        return out

    def text(self, day: date, shifts) -> str:
        """Shifts that start on ``day`` -> cell text (codes where they match; a leave by its note).

        Two days hold the same roster when their texts are equal, so a leave whose note
        is not a code's meaning (typed on the roster page) survives an untouched cell."""
        parts = []
        for kind, start, end, note in sorted((s[0], s[1], s[2], s[3] if len(s) > 3 else None) for s in shifts):
            local0, local1 = start.astimezone(self.admin.db.tz), end.astimezone(self.admin.db.tz)
            fits = [c for c in self.codes
                    if c.type == kind and c.start is not None and self.span(day, c.start, c.end) == (start, end)]
            match = next((c for c in fits if note and c.label == note), fits[0] if fits else None)
            if match:
                parts.append(match.code)
                continue
            rng = f"{local0.strftime('%H:%M')}-{local1.strftime('%H:%M')}"
            prefix = self.first_of_type.get(kind) if kind != "REGULAR" else None
            parts.append(f"{prefix.code} {rng}" if prefix else rng)
        return "/".join(parts)


# ---------------------------------------------------------------- dates

def _month_hint(text: str) -> date | None:
    """A month named in a sheet name or title: 2026-11, 11/2569, พ.ย. 2569, November 2026."""
    text = _text(text)
    m = re.search(r"(\d{4})\s*[-/.]\s*(\d{1,2})(?!\d)", text)
    if m and 1 <= int(m.group(2)) <= 12:
        return _year_month(int(m.group(1)), int(m.group(2)))
    m = re.search(r"(?<!\d)(\d{1,2})\s*[-/.]\s*(\d{4})", text)
    if m and 1 <= int(m.group(1)) <= 12:
        return _year_month(int(m.group(2)), int(m.group(1)))
    low = text.lower()
    for name in sorted(MONTHS, key=len, reverse=True):
        found = re.search(rf"\b{name}\b", low) if name.isascii() else name in low
        if found:
            year = re.search(r"(\d{4})", text)
            if year:
                return _year_month(int(year.group(1)), MONTHS[name])
    return None


def _year_month(year: int, month: int) -> date | None:
    year = year - 543 if year > 2400 else year  # พ.ศ.
    return date(year, month, 1) if 2000 <= year <= 2100 else None


def _header_date(value, hint: date | None, prev: date | None) -> date | None:
    """A date column header, or None when the cell is not a date."""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _text(value) if not _blank(value) else ""
    if not text:
        return None
    m = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})", text)
    if m:
        return _safe_date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    m = re.match(r"^(\d{1,2})\s*[/.-]\s*(\d{1,2})\s*[/.-]\s*(\d{2,4})$", text)
    if m:
        year = int(m.group(3))
        year = year + 2500 if year < 100 and year > 50 else year + 2000 if year < 100 else year
        return _safe_date(year - 543 if year > 2400 else year, int(m.group(2)), int(m.group(1)))
    m = re.match(r"^(\d{1,2})\s*[/.-]\s*(\d{1,2})$", text)
    if m and hint:
        return _safe_date(hint.year, int(m.group(2)), int(m.group(1)))
    m = re.match(r"^(\d{1,2})(?:\.0)?(?!\d)(?!\s*[:.]\d)", text)  # a day number, e.g. "1" or "1 (อา)"
    if m and hint and 1 <= int(m.group(1)) <= 31:
        day = int(m.group(1))
        if prev is None:
            base = hint
        elif day > prev.day:
            base = prev.replace(day=1)
        else:  # 16 ... 31, 1 ... 15: the numbers start again in the next month
            base = (prev.replace(day=28) + timedelta(days=4)).replace(day=1)
        return _safe_date(base.year, base.month, day)
    return None


def _safe_date(year: int, month: int, day: int) -> date | None:
    try:
        return date(year, month, day)
    except ValueError:
        return None


def _letter(index: int) -> str:
    from openpyxl.utils import get_column_letter

    return get_column_letter(index + 1)


# ------------------------------------------------------------- download

def build_roster(admin: HospitalAdmin, month: str | None = None, role: str | None = None) -> bytes:
    """One month's roster as an Excel file (people x days, cells in shift codes)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    first = _month_hint(month) if month else admin.db.today().replace(day=1)
    if first is None:
        raise StoreError("เดือนไม่ถูกต้อง (ใช้ YYYY-MM)")
    last = (first.replace(day=28) + timedelta(days=4)).replace(day=1) - timedelta(days=1)
    days = [first + timedelta(days=i) for i in range((last - first).days + 1)]
    codes = load_codes(admin)
    book = Codebook(admin, codes)
    people = [p for p in admin.staff()["staff"] if p["active"] and (not role or p["role"] == role)]
    order = {"SURGEON": 0, "ANESTHESIOLOGIST": 1, "SCRUB_NURSE": 2, "CIRCULATING_NURSE": 3}
    people.sort(key=lambda p: (order.get(p["role"], 9), p["name"]))
    held: dict[tuple[str, date], list] = {}
    with admin.store._transaction(write=False) as conn:
        note = ", note" if ("staff_schedules", "note") in admin._columns(conn) else ", NULL AS note"
        for r in conn.execute(
            f"""SELECT staff_id, shift_start, shift_end, shift_type{note} FROM staff_schedules
                WHERE shift_start >= %s AND shift_start < %s""",
            (admin._at(first, time()), admin._at(last + timedelta(days=1), time())),
        ):
            day = r["shift_start"].astimezone(admin.db.tz).date()
            held.setdefault((str(r["staff_id"]), day), []).append(
                (r["shift_type"] or "REGULAR", r["shift_start"], r["shift_end"], r["note"]))

    head_font, head_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="0E7490")
    weekend = PatternFill("solid", fgColor="F1F5F9")
    wb = Workbook()
    guide = wb.active
    guide.title = GUIDE
    for line in [
        f"ตารางเวร {MONTH_TH[first.month]} {first.year + 543} สำหรับ ORchestraAI",
        "1. แผ่นตารางเวร: แถว = บุคลากร (จับคู่ด้วยรหัสพนักงาน หรือชื่อถ้าไม่มีรหัส) · คอลัมน์ = วันที่",
        "2. ช่องหนึ่งใส่รหัสเวรของวันนั้น ใส่หลายเวรคั่นด้วย / เช่น ช/บ · ใส่เวลาเองได้ เช่น 07:30-15:30 หรือ ล 08:00-12:00 (ลาครึ่งวัน)",
        "3. ช่องว่าง = ไม่มีเวรวันนั้น · ทุกวันที่มีในไฟล์ ระบบแทนเวรเดิมของคนนั้นด้วยสิ่งที่อยู่ในช่อง",
        "   คนหรือวันที่ไม่มีในไฟล์ ระบบไม่แตะ (ลบแถวคนที่ไม่ต้องการแก้ออกได้)",
        f"4. แก้หรือเพิ่มรหัสเวรของโรงพยาบาลได้ที่แผ่น {CODES_SHEET} (เวลาเลิกน้อยกว่าเวลาเริ่ม = ข้ามคืน, 00:00-00:00 = ทั้งวัน)",
        "5. หัวคอลัมน์ใช้วันที่จริง หรือเลข 1-31 ก็ได้ถ้าชื่อแผ่นหรือหัวตารางบอกเดือน เช่น 2026-11 หรือ พ.ย. 2569",
        "6. อัปโหลดที่หน้า ข้อมูลโรงพยาบาล → เวรและวันลา ระบบตรวจและแสดงผลก่อน ยังไม่บันทึกจนกว่าจะกดยืนยัน",
    ]:
        guide.append([line])
    guide["A1"].font = Font(bold=True, size=14)
    guide.column_dimensions["A"].width = 120

    ws_codes = wb.create_sheet(CODES_SHEET)
    ws_codes.append(["รหัส", "ประเภท", "เริ่ม", "เลิก", "ความหมาย"])
    for c in codes:
        ws_codes.append(c.as_row())
    for cell in ws_codes[1]:
        cell.font, cell.fill = head_font, head_fill
    for letter, width in zip("ABCDE", (10, 14, 10, 10, 40)):
        ws_codes.column_dimensions[letter].width = width
        ws_codes.column_dimensions[letter].number_format = "@"
    for row in ws_codes.iter_rows(min_row=2):
        for cell in row:
            cell.number_format = "@"
    types = DataValidation(type="list", formula1='"' + ",".join(TYPES) + '"', allow_blank=True)
    types.showErrorMessage = False
    ws_codes.add_data_validation(types)
    types.add("B2:B200")

    ws = wb.create_sheet(first.strftime("%Y-%m"))
    ws.append([f"ตารางเวร {MONTH_TH[first.month]} {first.year + 543}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(["รหัสพนักงาน", "ชื่อ-นามสกุล", "บทบาท"] + days)
    ws.append(["", WEEKDAY_ROW, ""] + [DAY_TH[d.weekday()] for d in days])
    for cell in ws[2]:
        cell.font, cell.fill = head_font, head_fill
        cell.alignment = Alignment(horizontal="center")
    for cell in ws[2][3:]:
        cell.number_format = "d"
    for p in people:
        ws.append([p["code"], p["name"], p["role"]] + [book.text(d, held.get((p["id"], d), [])) or None for d in days])
    ws.freeze_panes = "D4"
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width, ws.column_dimensions["C"].width = 14, 26, 18
    last_row = max(ws.max_row, 300)
    for i, d in enumerate(days, start=3):
        letter = _letter(i)
        ws.column_dimensions[letter].width = 6.5
        ws.column_dimensions[letter].number_format = "@"
        for (cell,) in ws.iter_rows(min_row=3, max_row=ws.max_row, min_col=i + 1, max_col=i + 1):
            cell.number_format = "@"
            cell.alignment = Alignment(horizontal="center")
            if d.weekday() >= 5:
                cell.fill = weekend
    picks = DataValidation(type="list", formula1='"' + ",".join(c.code for c in codes)[:250] + '"', allow_blank=True)
    picks.showErrorMessage = False  # combinations (ช/บ) and times are fine too
    ws.add_data_validation(picks)
    picks.add(f"D4:{_letter(len(days) + 2)}{last_row}")
    for sheet in wb.worksheets:
        sheet.page_setup.orientation = "landscape"
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
    wb.active = 2
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# --------------------------------------------------------------- upload

@dataclass
class Grid:
    name: str
    first_row: int  # Excel row number of the header
    id_cols: dict  # "code" / "name" -> column index
    dates: list  # (column index, date)
    rows: list  # (excel row, cells)


def _roster_grid(name: str, grid: list[list], hint: date | None) -> Grid | str | None:
    """Find the header row of a roster sheet: an identity column and date columns."""
    labels = {"code": {"รหัสพนักงาน", "รหัส", "employeecode", "code"},
              "name": {"ชื่อ-นามสกุล", "ชื่อ", "ชื่อนามสกุล", "name", "บุคลากร"}}
    title_hint = hint
    for index, row in enumerate(grid[:10]):
        ids = {}
        for c, value in enumerate(row):
            key = _squash(_text(value)) if not _blank(value) else ""
            for kind, names in labels.items():
                if key in {_squash(n) for n in names} and kind not in ids:
                    ids[kind] = c
        if not ids:
            for value in row:  # a title row above the header may name the month
                if not _blank(value) and title_hint is None:
                    title_hint = _month_hint(_text(value))
            continue
        dates, prev = [], None
        for c, value in enumerate(row):
            if c in ids.values():
                continue
            d = _header_date(value, title_hint, prev)
            if d is not None:
                dates.append((c, d))
                prev = d
        if not dates:
            looks_like_days = any(re.match(r"^\d{1,2}(?:\.0)?$", _text(v)) for v in row if not _blank(v))
            return (f"แผ่น {name}: หัวคอลัมน์เป็นเลขวันที่ แต่ไม่รู้ว่าเป็นเดือนไหน "
                    "(ตั้งชื่อแผ่นเป็น 2026-11 หรือเลือกเดือนตอนอัปโหลด)") if looks_like_days else None
        rows = [(index + 2 + n, list(r)) for n, r in enumerate(grid[index + 1:]) if any(not _blank(v) for v in r)]
        return Grid(name, index + 1, ids, dates, rows)
    return None


def read_roster(filename: str, content: bytes, month: str | None = None) -> tuple[list[Code] | None, list[dict],
                                                                               list[Grid], list[str]]:
    """(codes from the file or None, code errors, roster grids, warnings)."""
    from .hospital_io import MAX_BYTES

    if not content:
        raise StoreError("ไฟล์ว่างเปล่า")
    if len(content) > MAX_BYTES:
        raise StoreError(f"ไฟล์ใหญ่เกิน {MAX_BYTES // (1024 * 1024)} MB")
    name = (filename or "").lower()
    if content[:2] == b"PK" or name.endswith(".xlsx"):
        sheets = _xlsx_grids(content)
    elif name.endswith(".xls"):
        raise StoreError("ไฟล์ .xls (Excel รุ่นเก่า) ยังไม่รองรับ กรุณาบันทึกเป็น .xlsx หรือ CSV UTF-8")
    else:
        sheets = [(filename or "CSV", _csv_grid(content))]
    fallback = _month_hint(month) if month else None
    codes, code_errors, grids, warnings = None, [], [], []
    for sheet_name, grid in sheets:
        if _squash(sheet_name) == _squash(GUIDE):
            continue
        if _squash(sheet_name) == _squash(CODES_SHEET):
            codes, code_errors = parse_codes(grid)
            continue
        found = _roster_grid(sheet_name, grid, _month_hint(sheet_name) or fallback)
        if isinstance(found, str):
            raise StoreError(found)
        if found is None:
            if any(any(not _blank(v) for v in r) for r in grid):
                warnings.append(f"ข้ามแผ่น {sheet_name}: ไม่พบหัวตารางเวร (คอลัมน์รหัสพนักงานหรือชื่อ และวันที่)")
            continue
        grids.append(found)
    if not grids:
        raise StoreError("ไม่พบตารางเวรในไฟล์นี้ (ใช้ไฟล์ที่ดาวน์โหลดจากหน้านี้เป็นแบบ)")
    return codes, code_errors, grids, warnings


def import_roster(admin: HospitalAdmin, filename: str, content: bytes, *, commit: bool, month: str | None = None,
                  actor: dict | None = None) -> dict:
    file_codes, errors, grids, warnings = read_roster(filename, content, month)
    codes = file_codes if file_codes else load_codes(admin)
    book = Codebook(admin, codes)
    today = admin.db.today()
    staff = admin.staff()["staff"]
    by_code = {p["code"].strip().lower(): p for p in staff}
    by_name: dict[str, list] = {}
    for p in staff:
        by_name.setdefault(_squash(p["name"]), []).append(p)

    wanted: dict[tuple[str, date], set] = {}  # (staff id, day) -> shifts from the file
    where: dict[tuple[str, date], str] = {}  # -> "sheet!cell" for messages
    names: dict[str, str] = {}
    all_days: set[date] = set()
    for g in grids:
        for col, day in g.dates:
            ref = f"{g.name}!{_letter(col)}{g.first_row}"
            if not today - timedelta(days=PAST_DAYS) <= day <= today + timedelta(days=FUTURE_DAYS):
                errors.append({"sheet": g.name, "cell": ref.split("!")[1],
                               "message": f"วันที่ {day.isoformat()} อยู่นอกช่วงที่แก้ได้ "
                                          f"({PAST_DAYS} วันย้อนหลังถึง {FUTURE_DAYS} วันข้างหน้า)"})
            all_days.add(day)
        if len({d for _, d in g.dates}) != len(g.dates):
            errors.append({"sheet": g.name, "cell": f"{g.first_row}", "message": "มีวันที่ซ้ำในหัวตาราง"})
        seen_people: dict[str, int] = {}
        for number, cells in g.rows:
            get = (lambda c: cells[c] if c is not None and c < len(cells) else None)  # noqa: E731
            code_cell, name_cell = get(g.id_cols.get("code")), get(g.id_cols.get("name"))
            if _blank(code_cell) and _blank(name_cell):
                continue
            if not _blank(name_cell) and _text(name_cell) == WEEKDAY_ROW:
                continue
            id_col = g.id_cols["code"] if "code" in g.id_cols and not _blank(code_cell) else g.id_cols.get("name",
                                                                                                  g.id_cols.get("code"))
            id_ref = f"{_letter(id_col)}{number}"
            person = None
            if not _blank(code_cell):
                person = by_code.get(_text(code_cell).lower())
                if person is None:
                    errors.append({"sheet": g.name, "cell": id_ref, "message": f"ไม่พบรหัสพนักงาน {_text(code_cell)}"})
                    continue
            else:
                matches = by_name.get(_squash(_text(name_cell)), [])
                if len(matches) != 1:
                    errors.append({"sheet": g.name, "cell": id_ref, "message":
                                   f"ไม่พบบุคลากรชื่อ {_text(name_cell)}" if not matches else
                                   f"มีบุคลากรชื่อ {_text(name_cell)} มากกว่า 1 คน (ใส่รหัสพนักงาน)"})
                    continue
                person = matches[0]
            if person["id"] in seen_people:
                errors.append({"sheet": g.name, "cell": id_ref,
                               "message": f"{person['name']} ซ้ำกับแถว {seen_people[person['id']]}"})
                continue
            seen_people[person["id"]] = number
            names[person["id"]] = person["name"]
            for col, day in g.dates:
                key = (person["id"], day)
                ref = f"{_letter(col)}{number}"
                if key in wanted:
                    errors.append({"sheet": g.name, "cell": ref, "message":
                                   f"{person['name']} วันที่ {day.isoformat()} มีในไฟล์มากกว่า 1 ที่ ({where[key]})"})
                    continue
                where[key] = f"{g.name}!{ref}"
                try:
                    wanted[key] = book.shifts(day, get(col)) if not _blank(get(col)) else set()
                except StoreError as exc:
                    errors.append({"sheet": g.name, "cell": ref, "message": str(exc)})
    if len(all_days) > MAX_DAYS:
        errors.append({"sheet": "", "cell": "", "message": f"ไฟล์มีวันที่เกิน {MAX_DAYS} วัน กรุณาแบ่งไฟล์"})
    if len(names) > MAX_PEOPLE:
        errors.append({"sheet": "", "cell": "", "message": f"มีบุคลากรเกิน {MAX_PEOPLE} คน"})

    report = {"filename": filename, "committed": False, "errors": errors, "warnings": warnings,
              "codes_from_file": file_codes is not None, "codes": [c.as_row() for c in codes],
              "from": min(all_days).isoformat() if all_days else None, "to": max(all_days).isoformat() if all_days else None,
              "people": len(names), "days": 0, "unchanged": 0, "added": 0, "removed": 0, "changes": [], "affected": []}
    if not wanted:
        return report

    first, last = min(all_days), max(all_days)
    with admin.store._transaction(write=commit and not errors) as conn:
        cols = admin._columns(conn)
        held: dict[tuple[str, date], set] = {}
        note = ", note" if ("staff_schedules", "note") in cols else ", NULL AS note"
        for r in conn.execute(
            f"""SELECT staff_id, shift_start, shift_end, shift_type{note} FROM staff_schedules
                WHERE shift_start >= %s AND shift_start < %s AND staff_id = ANY(%s)""",
            (admin._at(first, time()), admin._at(last + timedelta(days=1), time()),
             [uuid.UUID(i) for i in names]),
        ):
            key = (str(r["staff_id"]), r["shift_start"].astimezone(admin.db.tz).date())
            held.setdefault(key, set()).add((r["shift_type"] or "REGULAR", r["shift_start"], r["shift_end"], r["note"]))
        changed = {}
        for key, shifts in wanted.items():
            before = held.get(key, set())
            if book.text(key[1], shifts) == book.text(key[1], before):
                report["unchanged"] += 1
            else:
                changed[key] = (before, shifts)
        report["days"] = len(changed)
        report["removed"] = sum(len(b) for b, _ in changed.values())
        report["added"] = sum(len(a) for _, a in changed.values())
        per_person: dict[str, list] = {}
        for (sid, day), (before, after) in sorted(changed.items(), key=lambda kv: (names[kv[0][0]], kv[0][1])):
            per_person.setdefault(sid, []).append({"date": day.isoformat(), "before": book.text(day, before) or "-",
                                                   "after": book.text(day, after) or "-"})
        report["changes"] = [{"name": names[sid], "days": days} for sid, days in per_person.items()]
        if not commit or errors or not changed:
            return report
        with conn.cursor() as cur:
            for (sid, day) in changed:
                cur.execute("DELETE FROM staff_schedules WHERE staff_id = %s AND shift_start >= %s AND shift_start < %s",
                            (sid, admin._at(day, time()), admin._at(day + timedelta(days=1), time())))
            rows = [(sid, s[1], s[2], s[0], s[3]) for (sid, _), (_, after) in changed.items() for s in after]
            if rows:
                if ("staff_schedules", "note") in cols:
                    cur.executemany("INSERT INTO staff_schedules (staff_id, shift_start, shift_end, shift_type, note) "
                                    "VALUES (%s, %s, %s, %s, %s)", rows)
                else:
                    cur.executemany("INSERT INTO staff_schedules (staff_id, shift_start, shift_end, shift_type) "
                                    "VALUES (%s, %s, %s, %s)", [r[:4] for r in rows])
        if file_codes:
            _save_codes(admin, conn, file_codes)
        admin._log(conn, f"นำเข้าตารางเวรจาก {filename[:80]}: {len(per_person)} คน {len(changed)} วันที่เปลี่ยน",
                   [f"{first.isoformat()} ถึง {last.isoformat()} · เวรใหม่ {report['added']} · เวรเดิมที่ถูกแทน "
                    f"{report['removed']}"], actor, first)
        report["affected"] = admin._affected(conn, first, last + timedelta(days=1))
        report["committed"] = True
    return report
