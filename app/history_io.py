"""Past surgeries from Excel (or CSV), so the duration forecast learns from day one.

A row is a case done before the app was used: date, surgeon, procedure, and how long it took (time into
and out of the room, or the minutes). Room, urgency and the booked minutes are optional. Patient
columns (HN, name ...) are never read. Rows go into ``case_history`` (migration 014), which the
forecast (``Database.history``), the preference-card suggestions and the before/after report read.

Rows that would mislead the forecast are left out and listed: older than two years (the forecast
looks back two years), far from the procedure's standard time (data-entry slips), and cases already
imported. Wrong rows are listed with the reason; the right ones can be imported without them.
"""
from __future__ import annotations

import io
import re
import statistics
import uuid
from datetime import date, datetime, time, timedelta

from orchestra_engine.intake import mask_identifiers

from .hospital_io import MAX_BYTES, MAX_ROWS, _blank, _csv_grid, _text, _xlsx_grids, as_text, current
from .store import StoreError

SHEET = "ประวัติผ่าตัด"
GUIDE = "วิธีกรอก"
LISTS = "รายการ"
MAX_YEARS = 2
OUTLIER_LOW, OUTLIER_HIGH = 0.25, 4.0  # x the procedure's standard minutes
MAX_ERRORS = 200

# key, header in the template, other headers accepted (compared without spaces, case or text in brackets)
COLUMNS = [
    ("date", "วันที่ผ่าตัด", ("วันที่", "วันผ่าตัด", "date", "surgerydate", "opdate", "dateofsurgery")),
    ("surgeon", "ศัลยแพทย์ (รหัสหรือชื่อ)", ("ศัลยแพทย์", "แพทย์", "แพทย์ผู้ผ่าตัด", "รหัสแพทย์", "surgeon", "doctor")),
    ("procedure", "หัตถการ (ชื่อหรือรหัส ICD-9)", ("หัตถการ", "ชื่อหัตถการ", "รหัสหัตถการ", "procedure", "operation",
                                                   "icd9", "icd9cm")),
    ("in", "เวลาเข้าห้อง", ("เข้าห้อง", "เวลาเริ่ม", "in", "inroom", "timein", "start")),
    ("out", "เวลาออกห้อง", ("ออกห้อง", "เวลาเสร็จ", "out", "outroom", "timeout", "end")),
    ("minutes", "ระยะเวลาจริง (นาที)", ("ระยะเวลาจริง", "ระยะเวลา", "นาที", "duration", "actual", "actualmins")),
    ("booked", "เวลาที่จองไว้ (นาที)", ("เวลาที่จองไว้", "เวลาที่จอง", "booked", "planned", "bookedmins")),
    ("room", "ห้อง", ("ห้องผ่าตัด", "room", "or")),
    ("urgency", "ความเร่งด่วน", ("ประเภทเคส", "urgency")),
]
PATIENT = re.compile(r"^(hn|cid)$|ผู้ป่วย|patient|เลขบัตร|นามสกุล|ชื่อสกุล|อายุ|^age$")
URGENCY = {"elective": "ELECTIVE", "นัดล่วงหน้า": "ELECTIVE", "นัด": "ELECTIVE", "urgent": "URGENT", "ด่วน": "URGENT",
           "emergency": "EMERGENCY", "ฉุกเฉิน": "EMERGENCY"}
MAX_TOTAL_ROWS = 20000  # all sheets together (each sheet: MAX_ROWS)
BRACKETS = re.compile(r"\([^()]*\)")  # linear: "(.*?)" is slow on a cell full of "("
TITLES = re.compile(r"^(ศ\.|รศ\.|ผศ\.|นพ\.|พญ\.|ทพ\.|ทญ\.|dr\.?|prof\.?)+", re.IGNORECASE)


def _norm(text) -> str:
    """A header or name to compare: no bracketed text, spaces, punctuation or case."""
    text = BRACKETS.sub("", str(text or "")[:200])
    return re.sub(r"[\s\-_.:/]", "", text).lower()


def _person(text) -> str:
    text = BRACKETS.sub("", str(text or "")[:200]).strip()
    return re.sub(r"[\s.]", "", TITLES.sub("", text.replace(" ", ""))).lower()


HEADERS = {}
for _key, _header, _aliases in COLUMNS:
    for _h in (_header, *_aliases):
        HEADERS.setdefault(_norm(_h), _key)


# ------------------------------------------------------------------ cells

def parse_date(value) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and 20000 < value < 80000:  # an Excel day number
        return date(1899, 12, 30) + timedelta(days=int(value))
    text = _text(value).split(" ")[0]
    m = re.fullmatch(r"(\d{4})-(\d{1,2})-(\d{1,2})", text)
    if m:
        y, mo, d = (int(x) for x in m.groups())
    else:
        m = re.fullmatch(r"(\d{1,2})[/.\-](\d{1,2})[/.\-](\d{2}|\d{4})", text)
        if not m:
            raise ValueError(f"วันที่ไม่ถูกต้อง: {text} (ใช้ วว/ดด/ปปปป เช่น 31/10/2569)")
        d, mo, y = (int(x) for x in m.groups())
        if y < 100:  # two digits: 69 = 2569 (Thai year), 26 = 2026
            y = 2500 + y if y >= 40 else 2000 + y
    if y > 2400:  # Buddhist Era
        y -= 543
    try:
        return date(y, mo, d)
    except ValueError as exc:
        raise ValueError(f"วันที่ไม่ถูกต้อง: {text}") from exc


def parse_clock(value) -> int:
    """Minutes after midnight."""
    if isinstance(value, datetime):
        value = value.time()
    if isinstance(value, time):
        return value.hour * 60 + value.minute
    if isinstance(value, float) and 0 <= value < 1:  # Excel stores a time as a fraction of a day
        return round(value * 1440)
    text = _text(value).replace("น.", "").strip()
    m = re.fullmatch(r"(\d{1,2})[:.](\d{2})(?::\d{2})?", text) or re.fullmatch(r"(\d{2})(\d{2})", text)
    if not m or int(m.group(1)) > 23 or int(m.group(2)) > 59:
        raise ValueError(f"เวลาไม่ถูกต้อง: {text} (ใช้ ชม:นาที เช่น 08:30)")
    return int(m.group(1)) * 60 + int(m.group(2))


def _minutes(value, label: str) -> int:
    try:
        n = round(float(_text(value).replace(",", "")))
    except ValueError as exc:
        raise ValueError(f"{label}ต้องเป็นตัวเลขนาที: {_text(value)}") from exc
    if not 5 <= n <= 1440:
        raise ValueError(f"{label} {n} นาที ผิดปกติ (ต้องอยู่ระหว่าง 5 ถึง 1440)")
    return n


# ------------------------------------------------------------------ reading a file

def _table(grid: list[list]) -> tuple[int, dict[str, int], list[str]] | None:
    """(header row index, column of each key, patient columns left out) or None."""
    for i, row in enumerate(grid[:15]):
        cols, patient = {}, []
        for j, cell in enumerate(row):
            if _blank(cell):
                continue
            key = HEADERS.get(_norm(cell))
            if key and key not in cols:
                cols[key] = j
            elif PATIENT.search(_norm(cell)):
                patient.append(_text(cell))
        if {"date", "surgeon", "procedure"} <= set(cols):
            return i, cols, patient
    return None


def read_rows(filename: str, content: bytes) -> tuple[list[tuple[int, dict]], list[str]]:
    """((excel row number, {key: cell}) for each filled row, warnings)."""
    if not content:
        raise StoreError("ไฟล์ว่างเปล่า")
    if len(content) > MAX_BYTES:
        raise StoreError(f"ไฟล์ใหญ่เกิน {MAX_BYTES // (1024 * 1024)} MB")
    name = (filename or "").lower()
    if content[:2] == b"PK" or name.endswith(".xlsx"):
        sheets = _xlsx_grids(content)
    elif name.endswith(".xls"):
        raise StoreError("ไฟล์ .xls (Excel รุ่นเก่า) ยังไม่รองรับ กรุณาบันทึกเป็น .xlsx หรือ CSV UTF-8")
    else:
        sheets = [(filename or "CSV", _csv_grid(content))]
    rows, warnings = [], []
    for sheet, grid in sheets:
        if sheet in (GUIDE, LISTS):
            continue
        if len(grid) > MAX_ROWS + 10:
            raise StoreError(f"แผ่น {str(sheet)[:40]} มีเกิน {MAX_ROWS} แถว กรุณาแบ่งเป็นหลายไฟล์")
        found = _table(grid)
        if found is None:
            continue
        head, cols, patient = found
        if patient:
            warnings.append(f"ไม่นำเข้าคอลัมน์ข้อมูลผู้ป่วย ({', '.join(patient[:4])}) · ไม่จำเป็นต่อการประเมินเวลา")
        for i, raw in enumerate(grid[head + 1:], start=head + 2):
            cells = {k: raw[j] if j < len(raw) else None for k, j in cols.items()}
            if all(_blank(v) for v in cells.values()):
                continue
            rows.append((i, cells))
        if len(rows) > MAX_TOTAL_ROWS:
            raise StoreError(f"ไฟล์มีเกิน {MAX_TOTAL_ROWS} แถว กรุณาแบ่งเป็นหลายไฟล์")
    if not rows and not any(_table(g) for _, g in sheets):
        raise StoreError("ไม่พบหัวตาราง (ต้องมีคอลัมน์ วันที่ผ่าตัด, ศัลยแพทย์ และหัตถการ) · ใช้แบบฟอร์มที่ดาวน์โหลดจากหน้านี้")
    return rows, warnings


# ------------------------------------------------------------------ check and import

def _lookups(admin) -> tuple[dict, dict, dict, dict]:
    staff = [p for p in current(admin, "staff") if p["role"] == "SURGEON"]
    procs = current(admin, "procedures")
    surgeons = {**{_person(p["name"]): p for p in staff}, **{str(p["code"] or "").strip().lower(): p for p in staff if p["code"]}}
    by_code = {str(p["code"]).strip().lower(): p for p in procs if p.get("code")}
    by_name = {_norm(p["name"]): p for p in procs}
    return surgeons, by_code, by_name, {p["_id"]: p for p in procs}


def check(admin, rows: list[tuple[int, dict]], today: date) -> dict:
    """Every row read, matched and checked; nothing saved."""
    surgeons, by_code, by_name, procs = _lookups(admin)
    oldest = today - timedelta(days=365 * MAX_YEARS)
    good, errors, outliers, seen = [], [], [], set()
    old = duplicates = 0
    for number, cells in rows:
        try:
            missing = [label for key, label in (("date", "วันที่ผ่าตัด"), ("surgeon", "ศัลยแพทย์"), ("procedure", "หัตถการ"))
                       if _blank(cells.get(key))]
            if missing:
                raise ValueError("ต้องกรอก" + " ".join(missing))
            day = parse_date(cells["date"])
            if day > today:
                raise ValueError(f"วันที่ {day.isoformat()} ยังมาไม่ถึง")
            surgeon = surgeons.get(_text(cells["surgeon"]).lower()) or surgeons.get(_person(cells["surgeon"]))
            if surgeon is None:
                raise ValueError(f"ไม่พบศัลยแพทย์ {_text(cells['surgeon'])} (ใช้รหัสพนักงานหรือชื่อตามหน้าบุคลากร)")
            proc = by_code.get(_text(cells["procedure"]).lower()) or by_name.get(_norm(cells["procedure"]))
            if proc is None:
                raise ValueError(f"ไม่พบหัตถการ {_text(cells['procedure'])} (ใช้ชื่อหรือรหัส ICD-9 ตามหน้าหัตถการ)")
            start = parse_clock(cells["in"]) if not _blank(cells.get("in")) else None
            if start is not None and not _blank(cells.get("out")):
                actual = (parse_clock(cells["out"]) - start) % 1440
                if not 5 <= actual <= 1440:
                    raise ValueError(f"เวลาเข้า-ออกห้องห่างกัน {actual} นาที ผิดปกติ")
            elif not _blank(cells.get("minutes")):
                actual = _minutes(cells["minutes"], "ระยะเวลาจริง")
            else:
                raise ValueError("ใส่เวลาเข้าห้องและออกห้อง หรือระยะเวลาจริง (นาที)")
            booked = None if _blank(cells.get("booked")) else _minutes(cells["booked"], "เวลาที่จองไว้")
            room = None if _blank(cells.get("room")) else _text(cells["room"]).upper()[:20]
            urgency = "ELECTIVE"
            if not _blank(cells.get("urgency")):
                urgency = URGENCY.get(_text(cells["urgency"]).strip().lower())
                if urgency is None:
                    raise ValueError(f"ความเร่งด่วนไม่ถูกต้อง: {_text(cells['urgency'])} (นัดล่วงหน้า / ด่วน / ฉุกเฉิน)")
        except ValueError as exc:
            if len(errors) < MAX_ERRORS:
                errors.append({"row": number, "message": str(exc)})
            else:
                errors.append(None)
            continue
        if day < oldest:
            old += 1
            continue
        std = procs[proc["_id"]]["duration"] or 0
        if std and not OUTLIER_LOW * std <= actual <= OUTLIER_HIGH * std:
            outliers.append({"row": number, "text": f"{proc['name']} {actual} นาที (เวลามาตรฐาน {std} นาที)"})
            continue
        key = (day, surgeon["_id"], proc["_id"], room or "", start if start is not None else -1, actual)
        if key in seen:
            duplicates += 1
            continue
        seen.add(key)
        good.append({"row": number, "day": day, "start": start, "surgeon_id": surgeon["_id"], "surgeon": surgeon["name"],
                     "procedure_id": proc["_id"], "procedure": proc["name"], "room": room, "urgency": urgency,
                     "booked": booked, "actual": actual})
    return {"good": good, "errors": [e for e in errors if e], "error_count": len(errors), "old": old,
            "outliers": outliers, "duplicates": duplicates}


def _require(conn) -> None:
    if not conn.execute("SELECT to_regclass('public.case_history') IS NOT NULL AS ok").fetchone()["ok"]:
        raise StoreError("ยังนำเข้าประวัติไม่ได้: ให้รัน SQL ไฟล์ 014 (หน้า ⚙ สถานะระบบ → รัน SQL ที่ยังขาด) ก่อน")


def _stored_keys(conn, good: list[dict]) -> set:
    if not good:
        return set()
    first, last = min(g["day"] for g in good), max(g["day"] for g in good)
    return {(r["surgery_date"], str(r["surgeon_id"]), str(r["procedure_id"]), r["room_code"] or "",
             r["start_time"].hour * 60 + r["start_time"].minute if r["start_time"] else -1, r["actual_mins"])
            for r in conn.execute(
                """SELECT surgery_date, surgeon_id, procedure_id, room_code, start_time, actual_mins FROM case_history
                   WHERE surgery_date BETWEEN %s AND %s""", (first, last))}


def _key(g: dict) -> tuple:
    return (g["day"], g["surgeon_id"], g["procedure_id"], g["room"] or "", g["start"] if g["start"] is not None else -1,
            g["actual"])


def import_history(admin, filename: str, content: bytes, *, commit: bool, actor: dict | None = None) -> dict:
    """Check an uploaded file; with ``commit`` save its right rows (the wrong ones are listed, not saved)."""
    rows, warnings = read_rows(filename, content)
    filename = mask_identifiers(str(filename or "")[:200])[0]  # shown in the list and the activity log
    today = admin.db.today()
    result = check(admin, rows, today)
    good = result["good"]
    with admin.store._transaction(write=commit) as conn:
        _require(conn)
        stored = _stored_keys(conn, good)
        new = [g for g in good if _key(g) not in stored]
        imported, batch = 0, None
        if commit and new:
            batch = uuid.uuid4()
            who = None if actor is None else f"{actor.get('name')} ({actor.get('role')})"
            conn.execute("INSERT INTO case_history_imports (batch_id, filename, rows, imported_by) VALUES (%s, %s, 0, %s)",
                         (batch, (filename or "")[:200], who))
            for g in new:
                imported += conn.execute(
                    """INSERT INTO case_history (batch_id, surgery_date, start_time, surgeon_id, procedure_id, room_code,
                                                 urgency, booked_mins, actual_mins)
                       VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s) ON CONFLICT DO NOTHING""",
                    (batch, g["day"], time(g["start"] // 60, g["start"] % 60) if g["start"] is not None else None,
                     uuid.UUID(g["surgeon_id"]), uuid.UUID(g["procedure_id"]), g["room"], g["urgency"], g["booked"],
                     g["actual"])).rowcount
            conn.execute("UPDATE case_history_imports SET rows = %s WHERE batch_id = %s", (imported, batch))
            admin._log(conn, f"นำเข้าประวัติผ่าตัดย้อนหลัง {imported} เคส จาก {(filename or '')[:80]}",
                       [f"ข้ามแถวที่ผิด {result['error_count']}, เก่ากว่า {MAX_YEARS} ปี {result['old']}, "
                        f"เวลาผิดปกติ {len(result['outliers'])}, ซ้ำ {result['duplicates'] + len(good) - len(new)}"], actor)
    pairs: dict[tuple[str, str], list[int]] = {}
    for g in new:
        pairs.setdefault((g["surgeon"], g["procedure"]), []).append(g["actual"])
    summary = sorted(({"surgeon": s, "procedure": p, "n": len(xs), "median": round(statistics.median(xs))}
                      for (s, p), xs in pairs.items()), key=lambda x: (-x["n"], x["surgeon"], x["procedure"]))
    return {
        "filename": filename, "rows": len(rows), "ready": len(new), "errors": result["errors"],
        "error_count": result["error_count"], "old": result["old"], "outliers": result["outliers"][:100],
        "outlier_count": len(result["outliers"]), "duplicates": result["duplicates"], "already": len(good) - len(new),
        "warnings": warnings, "pairs": summary[:20], "pair_count": len(summary),
        "first": min((g["day"] for g in new), default=None), "last": max((g["day"] for g in new), default=None),
        "committed": bool(commit and imported), "imported": imported, "batch": str(batch) if batch else None,
        "max_years": MAX_YEARS,
    }


def overview(admin) -> dict:
    """What has been imported: totals and each file."""
    with admin.store._transaction(write=False) as conn:
        if not conn.execute("SELECT to_regclass('public.case_history') IS NOT NULL AS ok").fetchone()["ok"]:
            return {"ready": False, "batches": [], "cases": 0}
        totals = conn.execute(
            """SELECT count(*) AS cases, count(DISTINCT surgeon_id) AS surgeons, count(DISTINCT procedure_id) AS procedures,
                      min(surgery_date) AS first, max(surgery_date) AS last FROM case_history""").fetchone()
        batches = conn.execute(
            """SELECT b.batch_id, b.filename, b.imported_by, b.imported_at, count(h.history_id) AS rows
               FROM case_history_imports b LEFT JOIN case_history h ON h.batch_id = b.batch_id
               GROUP BY b.batch_id ORDER BY b.imported_at DESC""").fetchall()
    tz = admin.db.tz
    return {"ready": True, "cases": totals["cases"], "surgeons": totals["surgeons"], "procedures": totals["procedures"],
            "first": totals["first"].isoformat() if totals["first"] else None,
            "last": totals["last"].isoformat() if totals["last"] else None, "max_years": MAX_YEARS,
            "batches": [{"id": str(b["batch_id"]), "filename": b["filename"], "by": b["imported_by"], "rows": b["rows"],
                         "at": b["imported_at"].astimezone(tz).isoformat(timespec="minutes")} for b in batches]}


def delete_batch(admin, batch_id: str, *, actor: dict | None = None) -> dict:
    try:
        bid = uuid.UUID(batch_id)
    except ValueError as exc:
        raise StoreError("ไม่พบชุดที่นำเข้านี้") from exc
    with admin.store._transaction(write=True) as conn:
        _require(conn)
        row = conn.execute("SELECT filename FROM case_history_imports WHERE batch_id = %s", (bid,)).fetchone()
        if row is None:
            raise StoreError("ไม่พบชุดที่นำเข้านี้")
        removed = conn.execute("DELETE FROM case_history WHERE batch_id = %s", (bid,)).rowcount
        conn.execute("DELETE FROM case_history_imports WHERE batch_id = %s", (bid,))
        admin._log(conn, f"ลบประวัติผ่าตัดที่นำเข้าจาก {row['filename']}: {removed} เคส", [], actor)
    return {"removed": removed}


# ------------------------------------------------------------------ before / after report

def compare(conn, cases, first: date, last: date, tz) -> list[dict] | None:
    """Per procedure: actual minutes before the app (imported) and in the app in this period (medians)."""
    if not conn.execute("SELECT to_regclass('public.case_history') IS NOT NULL AS ok").fetchone()["ok"]:
        return None
    before: dict[str, list[int]] = {}
    for r in conn.execute(
            """SELECT p.name, h.actual_mins FROM case_history h JOIN procedures p ON p.procedure_id = h.procedure_id
               WHERE h.surgery_date > (NOW() - INTERVAL '2 years')::date"""):
        before.setdefault(r["name"], []).append(r["actual_mins"])
    now: dict[str, list[int]] = {}
    for c in cases:
        if c.demo or c.actual_mins is None or not first <= c.start.astimezone(tz).date() <= last:
            continue
        now.setdefault(c.procedure, []).append(c.actual_mins)
    out = []
    for name in set(before) | set(now):
        b, n = before.get(name, []), now.get(name, [])
        mb = round(statistics.median(b)) if b else None
        mn = round(statistics.median(n)) if n else None
        out.append({"procedure": name, "before_n": len(b), "before_median": mb, "now_n": len(n), "now_median": mn,
                    "change": mn - mb if mb is not None and mn is not None else None})
    out.sort(key=lambda x: (-(x["before_n"] + x["now_n"]), x["procedure"]))
    return out[:20]


# ------------------------------------------------------------------ template

def template(admin) -> bytes:
    """An Excel form: instructions, the history sheet with drop-down lists, and a hidden sheet of the lists."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    surgeons = sorted(p["name"] for p in current(admin, "staff") if p["role"] == "SURGEON")
    procs = sorted(p["name"] for p in current(admin, "procedures"))
    wb = Workbook()
    guide = wb.active
    guide.title = GUIDE
    lines = [
        ("ประวัติผ่าตัดย้อนหลัง (สำหรับให้ระบบประเมินเวลาผ่าตัด)", True),
        ("กรอกในแผ่น \"ประวัติผ่าตัด\" 1 แถว = 1 เคสที่ผ่าตัดไปแล้ว (ใช้ได้ย้อนหลังไม่เกิน 2 ปี)", False),
        ("ไม่ต้องใส่ HN หรือชื่อผู้ป่วย ระบบไม่เก็บข้อมูลผู้ป่วยจากไฟล์นี้", False),
        ("", False),
        ("คอลัมน์", True),
        ("วันที่ผ่าตัด: วว/ดด/ปปปป (พ.ศ. หรือ ค.ศ.) เช่น 31/10/2569 หรือ 2026-10-31 (ต้องกรอก)", False),
        ("ศัลยแพทย์: รหัสพนักงาน หรือชื่อตามหน้าบุคลากร (เลือกจากรายการได้) (ต้องกรอก)", False),
        ("หัตถการ: ชื่อ หรือรหัส ICD-9 ตามหน้าหัตถการ (เลือกจากรายการได้) (ต้องกรอก)", False),
        ("เวลาเข้าห้อง / เวลาออกห้อง: ชม:นาที เช่น 08:30 และ 10:15 หรือกรอก ระยะเวลาจริง (นาที) แทน", False),
        ("เวลาที่จองไว้ (นาที), ห้อง, ความเร่งด่วน (นัดล่วงหน้า / ด่วน / ฉุกเฉิน): ไม่บังคับ", False),
        ("", False),
        ("ระบบข้ามให้เอง: เคสเก่ากว่า 2 ปี, เคสที่เวลาต่างจากเวลามาตรฐานของหัตถการมาก (น่าจะพิมพ์ผิด), "
         "และเคสที่เคยนำเข้าแล้ว · ตรวจไฟล์ก่อนบันทึกทุกครั้ง", False),
        ("", False),
        ("ตัวอย่าง", True),
        ("31/10/2569 | " + (surgeons[0] if surgeons else "Dr.Somchai") + " | " + (procs[0] if procs else "Laparoscopic cholecystectomy")
         + " | 08:30 | 10:05 | | 90 | OR-1 | นัดล่วงหน้า", False),
    ]
    for text, bold in lines:
        guide.append([text])
        if bold:
            guide.cell(guide.max_row, 1).font = Font(bold=True)
    guide.column_dimensions["A"].width = 120

    ws = wb.create_sheet(SHEET)
    ws.append([header for _, header, _ in COLUMNS])
    fill = PatternFill("solid", fgColor="DDEBF7")
    for cell in ws[1]:
        cell.font, cell.fill = Font(bold=True), fill
    for letter, width in zip("ABCDEFGHI", (14, 28, 40, 13, 13, 18, 18, 10, 14)):
        ws.column_dimensions[letter].width = width
        for row in range(2, 2001):
            ws[f"{letter}{row}"].number_format = "@"  # keep 08:30 and 31/10/2569 as typed
    ws.freeze_panes = "A2"

    lists = wb.create_sheet(LISTS)
    for i, name in enumerate(surgeons, 1):
        lists.cell(i, 1, name)
    for i, name in enumerate(procs, 1):
        lists.cell(i, 2, name)
    for i, word in enumerate(("นัดล่วงหน้า", "ด่วน", "ฉุกเฉิน"), 1):
        lists.cell(i, 3, word)
    lists.sheet_state = "hidden"
    for col, letter, n in (("B", "A", len(surgeons)), ("C", "B", len(procs)), ("I", "C", 3)):
        if n:
            dv = DataValidation(type="list", formula1=f"='{LISTS}'!${letter}$1:${letter}${n}", allow_blank=True)
            dv.showErrorMessage = False  # codes and other spellings are accepted too; the import checks them
            ws.add_data_validation(dv)
            dv.add(f"{col}2:{col}2000")
    as_text(wb)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
