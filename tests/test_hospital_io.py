"""Excel import / export of hospital data, and clearing the demo data.

The parsing tests need no database; the rest run against a real PostgreSQL
server (TEST_DATABASE_URL), like tests/test_db.py.
"""
from __future__ import annotations

import io
from datetime import time

import pytest

from app import hospital_io as hio
from app.store import StoreError
from test_db import ADMIN_URL, _rows, _tomorrow, db_url, store  # noqa: F401  (fixtures)

ADMIN = {"name": "Admin A", "username": "a", "role": "admin", "staff_id": None}
needs_db = pytest.mark.skipif(not ADMIN_URL, reason="set TEST_DATABASE_URL to run database tests")


def col(sheet, key):
    return next(c for c in sheet.cols if c.key == key)


# ------------------------------------------------------------------ cells

def test_cells_accept_thai_words_and_excel_values():
    assert hio.parse_cell(col(hio.STAFF, "role"), "ศัลยแพทย์") == "SURGEON"
    assert hio.parse_cell(col(hio.STAFF, "role"), " scrub_nurse ") == "SCRUB_NURSE"
    assert hio.parse_cell(col(hio.STAFF, "active"), "ไม่") is False
    assert hio.parse_cell(col(hio.STAFF, "active"), True) is True  # an Excel TRUE cell
    assert hio.parse_cell(col(hio.ROOMS, "status"), "ซ่อมบำรุง") == "MAINTENANCE"
    clock = col(hio.ROOMS, "elective_start")
    assert [hio.parse_cell(clock, v) for v in ("8:00", "08.30", time(7, 45), 0.5)] == ["08:00", "08:30", "07:45", "12:00"]
    assert hio.parse_cell(col(hio.ROOMS, "capabilities"), "laminar_flow, ROBOTIC_READY;") == ["LAMINAR_FLOW", "ROBOTIC_READY"]
    assert hio.parse_cell(col(hio.CARDS, "equipment"), "LAP_TRAY x2; C_ARM\nLAP_TRAY ×1") == {"LAP_TRAY": 3, "C_ARM": 1}
    assert hio.parse_cell(col(hio.PROCEDURES, "skills"), "SCRUB_NURSE: robotic_certified; ANESTHESIOLOGIST: PEDS") == {
        "SCRUB_NURSE": ["ROBOTIC_CERTIFIED"], "ANESTHESIOLOGIST": ["PEDS"]}
    assert hio.parse_cell(col(hio.STAFF, "max_hours_per_day"), "9.5") == 9.5
    assert hio.parse_cell(col(hio.STAFF, "code"), 12.0) == "12"
    for c, bad, message in [(col(hio.STAFF, "role"), "พยาบาลประจำตึก", "บทบาท"),
                            (col(hio.STAFF, "active"), "maybe", "ใช่ หรือ ไม่"),
                            (col(hio.ROOMS, "base_turnover_mins"), "ครึ่งชั่วโมง", "ตัวเลข"),
                            (col(hio.ROOMS, "base_turnover_mins"), 30.5, "จำนวนเต็ม"),
                            (clock, "25:00", "HH:MM"),
                            (col(hio.PROCEDURES, "skills"), "ROBOTIC_CERTIFIED", "บทบาท: ทักษะ")]:
        with pytest.raises(hio._CellError, match=message):
            hio.parse_cell(c, bad)


def test_files_are_recognised_by_sheet_name_or_header():
    rows = "﻿รหัสพนักงาน *,ชื่อ-นามสกุล *,บทบาท *,สีโปรด\nD-1,นพ.ก,ศัลยแพทย์,ฟ้า\n,,,\n"
    [table], warnings = hio.read_tables("staff.csv", rows.encode("utf-8"))
    assert table.sheet.kind == "staff" and table.unknown == ["สีโปรด"] and len(table.rows) == 1
    assert table.rows[0][0] == 2  # Excel row number of the first data row
    thai_windows = "รหัสวัสดุ;ชื่อวัสดุ\nX-1;สกรู\n".encode("cp874")  # Excel's CSV on Thai Windows
    [table], _ = hio.read_tables("supplies.csv", thai_windows)
    assert table.sheet.kind == "supplies" and table.rows[0][1] == ["X-1", "สกรู"]
    titled = "รายชื่อห้อง ปี 2569\nรหัสห้อง,สถานะ\nOR-1,ใช้งานได้\n"  # a title row above the header
    [table], _ = hio.read_tables("rooms.csv", titled.encode())
    assert table.sheet.kind == "rooms" and table.rows == [(3, ["OR-1", "ใช้งานได้"])]
    for content, name, message in [(b"", "a.xlsx", "ว่าง"), (b"x" * (hio.MAX_BYTES + 1), "a.csv", "5 MB"),
                                   (b"\xd0\xcf\x11\xe0", "old.xls", ".xls"), (b"PK\x03\x04broken", "a.xlsx", "เปิดไฟล์"),
                                   ("ชื่อ,อายุ\nก,1\n".encode(), "x.csv", "ไม่พบข้อมูล")]:
        with pytest.raises(StoreError, match=message):
            hio.read_tables(name, content)


def test_a_zip_that_expands_too_far_is_refused():
    import zipfile

    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("xl/worksheets/sheet1.xml", b"0" * (hio.MAX_UNZIPPED + 1))
    with pytest.raises(StoreError, match="ใหญ่เกินไป"):
        hio.read_tables("bomb.xlsx", out.getvalue())


# --------------------------------------------------------- with a database

@pytest.fixture()
def admin(store):
    from app.admin_store import Inventory

    return Inventory(store)


def _book(store, wb) -> bytes:
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _open(data: bytes):
    from openpyxl import load_workbook

    return load_workbook(io.BytesIO(data))


def _row(ws, first_value):
    return next(r for r in ws.iter_rows(min_row=2) if r[0].value == first_value)


def _sheet(report, kind):
    return next(s for s in report["sheets"] if s["kind"] == kind)


@needs_db
def test_downloaded_file_reads_back_unchanged(admin):
    data = hio.build_workbook(admin)
    wb = _open(data)
    assert wb.sheetnames == [hio.GUIDE] + [s.title for s in hio.SHEETS]
    staff = wb["บุคลากร"]
    assert staff["A1"].value == "รหัสพนักงาน *" and staff.column_dimensions["A"].number_format == "@"
    assert any(dv.formula1.startswith('"SURGEON') for dv in staff.data_validations.dataValidation)
    supplies = wb["วัสดุและรากเทียม"]
    [header] = supplies.iter_rows(max_row=1, values_only=True)
    first = next(supplies.iter_rows(min_row=2, values_only=True))
    assert first[header.index("คงคลังในระบบ (ดูอย่างเดียว)")] is not None
    assert first[header.index("นับได้จริง (กรอกเมื่อนับสต็อก)")] is None  # left blank: no stale stock counts

    report = hio.import_file(admin, "same.xlsx", data, commit=True)
    assert report["errors"] == 0 and report["new"] == report["changed"] == 0 and report["unchanged"] > 50
    assert not report["committed"]  # nothing to save
    blank = _open(hio.build_workbook(admin, blank=True))
    assert blank["ห้องผ่าตัด"].max_row == 1


@needs_db
def test_import_adds_and_changes_records_after_a_preview(admin, store, db_url):
    from test_db import _request

    day = _tomorrow(store)  # a booking that the room change below will break
    assert store.slot_check(_request(store.state(day), "Dr.Napat", "Laparoscopic cholecystectomy", "HN-IO-1"),
                            "09:00", commit=True, day=day)["booked"]
    room_code = store.state(day)["schedule"]["bookings"][0]["or_id"]
    wb = _open(hio.build_workbook(admin))
    _row(wb["ห้องผ่าตัด"], room_code)[5].value = "ซ่อมบำรุง"
    wb["บุคลากร"].append(["D-9001", "นพ.นำเข้า ทดสอบ", "ศัลยแพทย์", "Ortho", "robotic_certified", 40, 9, 10, "ใช่"])
    wb["วัสดุและรากเทียม"].append(["SCREW-7", "สกรู 7 มม.", "รากเทียม", "ตัว", None, 20, 5, "ใช่"])
    wb["หัตถการ"].append(["81.54", "Total knee replacement (import)", 150, "สะอาด", "ไม่", "ใช่", None, "LAMINAR_FLOW",
                          "SCRUB_NURSE: ORTHO_TRAINED", "screw-7 x2"])
    wb["Preference card"].append(["d-9001", None, "81.54", 140, 20, 0, 2, 1, 0, "C_ARM; ORTHO_KNEE_SET x2", "ท่านอนหงาย"])
    data = _book(store, wb)
    activity_before = len(store.state()["activity"])

    preview = hio.import_file(admin, "edit.xlsx", data, commit=False, actor=ADMIN)
    assert (preview["new"], preview["changed"], preview["errors"], preview["committed"]) == (4, 1, 0, False)
    [room] = _sheet(preview, "rooms")["rows"]
    assert room["key"] == room_code and room["changes"] == ["สถานะ: AVAILABLE → MAINTENANCE"]
    assert all(p["code"] != "D-9001" for p in admin.staff()["staff"])  # the preview saved nothing

    done = hio.import_file(admin, "edit.xlsx", data, commit=True, actor=ADMIN)
    assert done["committed"] and day in [a["date"] for a in done["affected"]]  # that booking's room is now closed
    person = next(p for p in admin.staff()["staff"] if p["code"] == "D-9001")
    assert (person["role"], person["skills"], person["max_hours_per_day"]) == ("SURGEON", ["ROBOTIC_CERTIFIED"], 9)
    item = next(i for i in admin.inventory()["items"] if i["code"] == "SCREW-7")
    assert (item["category"], item["on_hand"], item["unit"]) == ("IMPLANT", 20, "ตัว")
    proc = next(p for p in admin.procedures()["procedures"] if p["code"] == "81.54")
    assert proc["skills"] == {"SCRUB_NURSE": ["ORTHO_TRAINED"]} and proc["caps"] == ["LAMINAR_FLOW"] and proc["blood"]
    assert next(p for p in admin.inventory()["procedures"] if p["id"] == proc["id"])["supplies"] == {"SCREW-7": 2}
    card = next(c for c in admin.cards()["cards"] if c["procedure_id"] == proc["id"])
    assert card["surgeon"] == "นพ.นำเข้า ทดสอบ" and card["notes"] == "ท่านอนหงาย"
    assert card["roles"] == {"ANESTHESIOLOGIST": 0, "SCRUB_NURSE": 2, "CIRCULATING_NURSE": 1, "SURGEON": 0}
    assert sorted((e["name"], e["count"]) for e in card["equipment"]) == [("C_ARM", 1), ("ORTHO_KNEE_SET", 2)]
    activity = store.state()["activity"]
    assert len(activity) == activity_before + 1  # one entry for the whole file, not one per row
    assert activity[0]["title"].startswith("นำเข้าข้อมูลจาก edit.xlsx: เพิ่ม 4 แก้ไข 1")

    again = hio.import_file(admin, "edit.xlsx", hio.build_workbook(admin), commit=False)
    assert again["new"] == again["changed"] == again["errors"] == 0
    same = hio.import_file(admin, "edit.xlsx", data, commit=False)  # card row names its procedure by ICD code
    assert _sheet(same, "cards")["changed"] == 0 and same["changed"] == 0


@needs_db
def test_any_wrong_row_blocks_the_whole_file(admin, store):
    wb = _open(hio.build_workbook(admin, blank=True))
    wb["บุคลากร"].append(["N-1", "พยาบาลดี", "SCRUB_NURSE"])
    wb["บุคลากร"].append(["N-2", "พยาบาลผิด", "พยาบาลประจำตึก"])
    wb["บุคลากร"].append(["N-1", "ซ้ำ", "SCRUB_NURSE"])
    wb["บุคลากร"].append(["N-3", None, "SCRUB_NURSE"])
    wb["ห้องผ่าตัด"].append(["OR-9", "STANDARD", None, 30, 60, "AVAILABLE", "16:00", "08:00"])
    wb["Preference card"].append(["SCR-1", None, "Laparoscopic cholecystectomy", 60])
    wb["Preference card"].append(["SUR-GEN", None, "ไม่มีหัตถการนี้", 60])
    wb["เครื่องมือรายชิ้น"].append(["X-1", "ไม่มีชนิดนี้"])
    report = hio.import_file(admin, "bad.xlsx", _book(store, wb), commit=True)
    assert not report["committed"]
    errors = {(s["kind"], e["row"]): e["message"] for s in report["sheets"] for e in s["errors"]}
    assert "บทบาท" in errors[("staff", 3)] and "ซ้ำกับแถว 2" in errors[("staff", 4)]
    assert "ชื่อ-นามสกุล" in errors[("staff", 5)]
    assert "หลังเวลาเปิด" in errors[("rooms", 2)]  # the save code's own check
    assert "ไม่ใช่ศัลยแพทย์" in errors[("cards", 2)] and "ไม่พบหัตถการ" in errors[("cards", 3)]
    assert "ไม่พบชนิดเครื่องมือ" in errors[("units", 2)]
    assert report["new"] == 1  # N-1 alone was fine ...
    assert all(p["code"] != "N-1" for p in admin.staff()["staff"])  # ... but nothing was saved


@needs_db
def test_left_out_columns_keep_values_and_stock_is_counted_only_when_typed(admin, store):
    from openpyxl import Workbook

    person = next(p for p in admin.staff()["staff"] if p["code"] == "SCR-1")
    wb = Workbook()
    ws = wb.active
    ws.title = "บุคลากร"
    ws.append(["รหัสพนักงาน", "ชื่อ-นามสกุล", "ความเชี่ยวชาญ", "ชั่วโมงสูงสุดต่อวัน"])
    ws.append(["scr-1", "Malee S.", None, None])  # matched case-insensitively; blanks: text cleared, number kept
    report = hio.import_file(admin, "names.xlsx", _book(store, wb), commit=True)
    assert report["committed"] and report["changed"] == 1
    after = next(p for p in admin.staff()["staff"] if p["code"] == "SCR-1")
    assert (after["name"], after["role"], after["skills"]) == ("Malee S.", person["role"], person["skills"])
    assert after["max_hours_per_day"] == person["max_hours_per_day"]

    data = hio.build_workbook(admin)  # an old download ...
    code = admin.inventory()["items"][0]["code"]
    admin.count_stock(code, {"on_hand": 1})  # ... and stock moved since
    assert hio.import_file(admin, "old.xlsx", data, commit=True)["changed"] == 0  # stock not put back
    wb = _open(data)
    _row(wb["วัสดุและรากเทียม"], code)[5].value = 7  # นับได้จริง
    report = hio.import_file(admin, "count.xlsx", _book(store, wb), commit=True)
    assert report["committed"] and _sheet(report, "supplies")["rows"][0]["changes"] == [
        "นับได้จริง (กรอกเมื่อนับสต็อก): 1 → 7"]
    assert next(i for i in admin.inventory()["items"] if i["code"] == code)["on_hand"] == 7


@needs_db
def test_one_sheet_as_csv(admin):
    csv_text = "รหัสห้อง,ประเภทห้อง,ความสามารถของห้อง,เริ่มรับเคส elective,ปิดรับเคส elective\n" \
               "OR-7,ห้องความดันลบ,\"LAMINAR_FLOW, NEGATIVE_PRESSURE\",07.30,15:30\n"
    report = hio.import_file(admin, "rooms.csv", csv_text.encode("utf-8-sig"), commit=True, actor=ADMIN)
    assert report["committed"] and report["new"] == 1
    room = next(r for r in admin.rooms()["rooms"] if r["code"] == "OR-7")
    assert room["type"] == "NEGATIVE_PRESSURE" and room["capabilities"] == ["LAMINAR_FLOW", "NEGATIVE_PRESSURE"]
    assert (room["elective_start"], room["elective_end"], room["base_turnover_mins"]) == ("07:30", "15:30", 30)


@needs_db
def test_demo_cases_go_and_real_cases_stay(admin, store):
    from test_db import _request

    day = _tomorrow(store)
    request = _request(store.state(day), "Dr.Napat", "Laparoscopic cholecystectomy", "HN-REAL-1")
    assert store.slot_check(request, "09:00", commit=True, day=day)["booked"]
    status = hio.data_status(admin)
    assert status["demo_cases"] == 10 and status["demo_history"] > 50

    removed = hio.delete_demo_cases(admin, actor=ADMIN)
    assert removed["requests"] == status["demo_cases"] + status["demo_history"]
    assert store.state()["schedule"]["bookings"] == []  # today's demo day is gone
    assert [b["case_id"] for b in store.state(day)["schedule"]["bookings"]]  # the real case stays
    assert hio.data_status(admin)["demo_cases"] == 0
    assert store.report()["demo_cases"] == 0


@needs_db
def test_wipe_keeps_user_accounts_and_the_file_rebuilds_the_hospital(admin, store, db_url):
    data = hio.build_workbook(admin)
    with __import__("psycopg").connect(db_url, autocommit=True) as conn:
        conn.execute("""INSERT INTO app_users (username, display_name, role, password_hash, staff_id)
                        SELECT 'doc', 'Doc', 'surgeon', 'x', staff_id FROM staff WHERE employee_code = 'SUR-GEN'""")
    with pytest.raises(StoreError, match="ล้างข้อมูล"):
        hio.wipe_hospital(admin, "yes")
    done = hio.wipe_hospital(admin, " ล้างข้อมูล ", actor=ADMIN)
    assert done["deleted"]["บุคลากร"] > 10 and done["unlinked_users"] == 1
    assert all(n == 0 for n in hio.data_status(admin)["counts"].values())
    assert _rows(db_url, "SELECT username, staff_id FROM app_users") == [("doc", None)]
    assert store.state()["empty"]

    report = hio.import_file(admin, "backup.xlsx", data, commit=True, actor=ADMIN)
    assert report["committed"] and report["errors"] == 0 and report["changed"] == 0
    assert len(admin.cards()["cards"]) == _sheet(report, "cards")["new"] > 0  # cards found staff from the same file
    assert len(admin.equipment()["units"]) == _sheet(report, "units")["new"]
