"""Monthly rosters in Excel: shift codes, headers, preview and import.

The parsing tests need no database; the rest run against a real PostgreSQL
server (TEST_DATABASE_URL), like tests/test_db.py.
"""
from __future__ import annotations

import io
from datetime import date, datetime, time, timedelta
from types import SimpleNamespace
from zoneinfo import ZoneInfo

import pytest

from app import roster_io as rio
from app.store import StoreError
from test_db import ADMIN_URL, _rows, _tomorrow, db_url, store  # noqa: F401  (fixtures)

ADMIN = {"name": "Admin A", "username": "a", "role": "admin", "staff_id": None}
TZ = ZoneInfo("Asia/Bangkok")
needs_db = pytest.mark.skipif(not ADMIN_URL, reason="set TEST_DATABASE_URL to run database tests")


def at(day: date, clock: str, days: int = 0) -> datetime:
    return datetime.combine(day + timedelta(days=days), time.fromisoformat(clock), tzinfo=TZ)


@pytest.fixture()
def book():
    fake = SimpleNamespace(db=SimpleNamespace(tz=TZ), _at=lambda d, t: datetime.combine(d, t, tzinfo=TZ))
    return rio.Codebook(fake, list(rio.DEFAULT_CODES))


# ------------------------------------------------------------- no database

def test_shift_codes_and_times_in_a_cell(book):
    day = date(2026, 11, 3)
    assert book.shifts(day, "ช/บ") == {("REGULAR", at(day, "08:00"), at(day, "16:00"), None),
                                       ("REGULAR", at(day, "16:00"), at(day, "00:00", 1), None)}
    assert book.shifts(day, "oc") == {("ON_CALL", at(day, "16:00"), at(day, "08:00", 1), None)}
    assert book.shifts(day, "ลป") == {("LEAVE", at(day, "00:00"), at(day, "00:00", 1), "ลาป่วย")}
    assert book.shifts(day, " - ") == set() and book.shifts(day, "หยุด") == set()
    assert book.shifts(day, "07.30-15.30") == {("REGULAR", at(day, "07:30"), at(day, "15:30"), None)}
    assert book.shifts(day, "ล 08:00-12:00 + ช") == {("LEAVE", at(day, "08:00"), at(day, "12:00"), None),
                                                     ("REGULAR", at(day, "08:00"), at(day, "16:00"), None)}
    for bad in ("XX", "ช/XX", "OFF 08:00-12:00", "25:00-26:00"):
        with pytest.raises(StoreError):
            book.shifts(day, bad)
    # back into cell text: codes where they fit, a leave by its kind, other times as ranges
    assert book.text(day, book.shifts(day, "ช/บ")) == "ช/บ"
    assert book.text(day, book.shifts(day, "ลพ")) == "ลพ"
    assert book.text(day, {("LEAVE", at(day, "00:00"), at(day, "00:00", 1), "ไปประชุมต่างประเทศ")}) == "ล"
    assert book.text(day, {("ON_CALL", at(day, "08:00"), at(day, "23:59"), None)}) == "OC 08:00-23:59"
    assert book.text(day, set()) == ""


def test_hospital_codes_sheet():
    codes, errors = rio.parse_codes([["รหัสเวร"], ["รหัส", "ประเภท", "เริ่ม", "เลิก", "ความหมาย"],
                                     ["M", "เวรปกติ", "07:00", "15:00", "เช้า"], ["N", "regular", "23.00", "07.00"],
                                     ["V", "ลา", time(0), "24:00", "พักร้อน"], ["x", "หยุด"],
                                     ["M", "REGULAR", "07:00", "15:00"], ["B", "ทำงาน", "08:00", "16:00"],
                                     ["C", "REGULAR", "08:00"], ["a/b", "REGULAR", "08:00", "16:00"]])
    assert [(c.code, c.type, c.start, c.end) for c in codes] == [
        ("M", "REGULAR", time(7), time(15)), ("N", "REGULAR", time(23), time(7)), ("V", "LEAVE", time(0), time(0)),
        ("x", "OFF", None, None)]
    assert [e["cell"] for e in errors] == ["A7", "A8", "A9", "A10"]
    for e, text in zip(errors, ("รหัส M ซ้ำกับแถว 3", "ประเภทไม่ถูกต้อง", "ต้องใส่เวลาเริ่มและเวลาเลิก", "ห้ามมี /")):
        assert text in e["message"]


def test_date_headers_and_month_names():
    assert rio._month_hint("ตารางเวร พฤศจิกายน 2569") == date(2026, 11, 1)
    assert rio._month_hint("พ.ย. 69") is None and rio._month_hint("Summary") is None
    assert [rio._month_hint(t) for t in ("2026-11", "11/2569", "Nov 2026")] == [date(2026, 11, 1)] * 3
    hint, prev, days = date(2026, 11, 1), None, []
    for cell in (16, "30", "1 (อา)", 2.0):  # a roster from the 16th into the next month
        prev = rio._header_date(cell, hint, prev)
        days.append(prev)
    assert days == [date(2026, 11, 16), date(2026, 11, 30), date(2026, 12, 1), date(2026, 12, 2)]
    assert rio._header_date("1/11/2569", None, None) == rio._header_date(datetime(2026, 11, 1), None, None)
    assert rio._header_date("2026-11-03", None, None) == date(2026, 11, 3)
    assert rio._header_date("3", None, None) is None  # a day number needs a month
    assert rio._header_date("บทบาท", hint, None) is None and rio._header_date("12:00", hint, None) is None
    with pytest.raises(StoreError, match="ไม่รู้ว่าเป็นเดือนไหน"):
        rio.read_roster("roster.csv", "ชื่อ,1,2,3\nนก,ช,ช,ช\n".encode())


# ------------------------------------------------------------ with database

@pytest.fixture()
def admin(store):
    from app.admin_store import Inventory

    return Inventory(store)


def _open(data: bytes):
    from openpyxl import load_workbook

    return load_workbook(io.BytesIO(data))


def _save(wb) -> bytes:
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _row(ws, code):
    return next(i for i in range(4, ws.max_row + 1) if ws.cell(i, 1).value == code)


def _col(ws, day: date):
    return next(c for c in range(4, ws.max_column + 1) if ws.cell(2, c).value.date() == day)


def _shifts(db_url, code, day):
    return _rows(db_url, """SELECT s.shift_type, to_char(s.shift_start AT TIME ZONE 'Asia/Bangkok', 'HH24:MI'),
                                   to_char(s.shift_end AT TIME ZONE 'Asia/Bangkok', 'HH24:MI'), s.note
                            FROM staff_schedules s JOIN staff p USING (staff_id)
                            WHERE p.employee_code = %s AND (s.shift_start AT TIME ZONE 'Asia/Bangkok')::date = %s
                            ORDER BY s.shift_start""", code, day)


@needs_db
def test_downloaded_roster_reads_back_unchanged(admin):
    month = admin.db.today().strftime("%Y-%m")
    data = rio.build_roster(admin, month)
    wb = _open(data)
    assert wb.sheetnames == [rio.GUIDE, rio.CODES_SHEET, month]
    ws = wb[month]
    assert [ws.cell(2, c).value for c in (1, 2, 3)] == ["รหัสพนักงาน", "ชื่อ-นามสกุล", "บทบาท"]
    assert ws.cell(3, 2).value == rio.WEEKDAY_ROW and isinstance(ws.cell(2, 4).value, datetime)
    assert ws.cell(_row(ws, "SUR-GEN"), 4).value == "ช"  # the demo's 08:00-16:00 is the morning code
    report = rio.import_roster(admin, "same.xlsx", data, commit=True)
    assert report["errors"] == [] and report["days"] == 0 and not report["committed"]
    assert report["unchanged"] == report["people"] * len([c for c in range(4, ws.max_column + 1)])
    surgeons = _open(rio.build_roster(admin, month, "SURGEON"))[month]
    assert {surgeons.cell(r, 3).value for r in range(4, surgeons.max_row + 1)} == {"SURGEON"}


@needs_db
def test_roster_import_replaces_the_days_in_the_file(admin, store, db_url):
    from test_db import _request

    day = date.fromisoformat(_tomorrow(store))
    assert store.slot_check(_request(store.state(day.isoformat()), "Dr.Napat", "Laparoscopic cholecystectomy",
                                     "HN-RX-1"), "09:00", commit=True, day=day.isoformat())["booked"]
    month = day.strftime("%Y-%m")
    wb = _open(rio.build_roster(admin, month))
    ws = wb[month]
    napat, malee, col = _row(ws, "SUR-GEN"), _row(ws, "SCR-1"), _col(ws, day)
    ws.cell(napat, col).value = "ลพ"  # on leave the day of a booked case
    ws.cell(malee, col).value = "ช/บ"
    if col + 1 <= ws.max_column:
        ws.cell(malee, col + 1).value = None  # blank: no shift that day
    ws.cell(_row(ws, "SCR-2"), col).value = "07:30-15:30"
    data = _save(wb)
    logged_on = day.replace(day=1).isoformat()  # the import is logged on the first day the file covers
    activity = len(store.state(logged_on)["activity"])

    preview = rio.import_roster(admin, "nov.xlsx", data, commit=False)
    assert preview["errors"] == [] and not preview["committed"]
    changes = {c["name"]: c["days"] for c in preview["changes"]}
    assert changes["Dr.Napat (General)"] == [{"date": day.isoformat(), "before": "ช", "after": "ลพ"}]
    assert _shifts(db_url, "SUR-GEN", day) == [("REGULAR", "08:00", "16:00", None)]  # nothing saved yet

    done = rio.import_roster(admin, "nov.xlsx", data, commit=True, actor=ADMIN)
    assert done["committed"] and day.isoformat() in [a["date"] for a in done["affected"]]
    assert _shifts(db_url, "SUR-GEN", day) == [("LEAVE", "00:00", "00:00", "ลาพักร้อน")]
    assert _shifts(db_url, "SCR-1", day) == [("REGULAR", "08:00", "16:00", None), ("REGULAR", "16:00", "00:00", None)]
    assert _shifts(db_url, "SCR-2", day) == [("REGULAR", "07:30", "15:30", None)]
    if col + 1 <= ws.max_column:
        assert _shifts(db_url, "SCR-1", day + timedelta(days=1)) == []
    assert _shifts(db_url, "SCR-3", day) == [("REGULAR", "08:00", "16:00", None)]  # untouched cell: kept
    assert len(store.state(logged_on)["activity"]) == activity + 1  # one entry for the file
    assert rio.import_roster(admin, "nov.xlsx", data, commit=False)["days"] == 0


@needs_db
def test_untouched_cells_keep_leave_notes_and_people_not_in_the_file(admin, store, db_url):
    day = admin.db.today() + timedelta(days=3)
    who = {p["code"]: p["id"] for p in admin.staff()["staff"]}
    admin.add_shifts({"type": "LEAVE", "date_from": day.isoformat(), "date_to": day.isoformat(), "start": "00:00",
                      "end": "00:00", "staff_ids": [who["SCR-1"]], "note": "ไปประชุมต่างประเทศ"})
    month = day.strftime("%Y-%m")
    wb = _open(rio.build_roster(admin, month))
    ws = wb[month]
    assert ws.cell(_row(ws, "SCR-1"), _col(ws, day)).value == "ล/ช"
    for r in range(ws.max_row, 3, -1):  # keep only two people in the file
        if ws.cell(r, 1).value not in ("SCR-1", "SCR-2"):
            ws.delete_rows(r)
    ws.cell(_row(ws, "SCR-2"), _col(ws, day)).value = "ล"
    report = rio.import_roster(admin, "two.xlsx", _save(wb), commit=True)
    assert report["committed"] and report["people"] == 2 and report["days"] == 1
    assert ("LEAVE", "00:00", "00:00", "ไปประชุมต่างประเทศ") in _shifts(db_url, "SCR-1", day)
    assert _shifts(db_url, "SCR-2", day) == [("LEAVE", "00:00", "00:00", "ลา (ทั้งวัน)")]
    assert _shifts(db_url, "SUR-GEN", day) == [("REGULAR", "08:00", "16:00", None)]  # not in the file


@needs_db
def test_wrong_cells_block_the_import(admin):
    month = admin.db.today().strftime("%Y-%m")
    wb = _open(rio.build_roster(admin, month))
    ws = wb[month]
    ws.cell(_row(ws, "SCR-1"), 5).value = "ช+XX"
    ws.cell(_row(ws, "SCR-2"), 1).value = "NOBODY"
    ws.append(["SCR-3", "again", None, "ช"])
    ws.append([None, "ไม่มีชื่อนี้", None, "ช"])
    ws.cell(_row(ws, "SCR-4"), 4).value = "ด"
    report = rio.import_roster(admin, "bad.xlsx", _save(wb), commit=True)
    assert not report["committed"]
    messages = {e["cell"]: e["message"] for e in report["errors"]}
    assert "XX" in messages[f"E{_row(ws, 'SCR-1')}"]
    assert "ไม่พบรหัสพนักงาน NOBODY" in messages[f"A{_row(ws, 'NOBODY')}"]
    assert "ซ้ำกับแถว" in messages[f"A{ws.max_row - 1}"] and "ไม่พบบุคลากรชื่อ" in messages[f"B{ws.max_row}"]
    assert report["days"] >= 1  # the good change is shown, but nothing was saved
    assert rio.import_roster(admin, "bad.xlsx", _save(wb), commit=False)["days"] == report["days"]

    far = _open(rio.build_roster(admin, (admin.db.today() + timedelta(days=500)).strftime("%Y-%m")))
    report = rio.import_roster(admin, "far.xlsx", _save(far), commit=False)
    assert report["errors"] and "นอกช่วง" in report["errors"][0]["message"]


@needs_db
def test_own_codes_names_and_day_number_headers(admin, db_url):
    day = admin.db.today() + timedelta(days=2)
    month = day.strftime("%Y-%m")
    wb = _open(rio.build_roster(admin, month))
    wb[rio.CODES_SHEET].append(["ชบ", "REGULAR", "08:00", "00:00", "เช้าต่อบ่าย"])
    ws = wb[month]
    ws.cell(_row(ws, "SCR-1"), _col(ws, day)).value = "ชบ"
    assert rio.import_roster(admin, "codes.xlsx", _save(wb), commit=True)["committed"]
    assert _shifts(db_url, "SCR-1", day) == [("REGULAR", "08:00", "00:00", None)]
    again = _open(rio.build_roster(admin, month))  # the hospital's codes are kept for the next download
    assert "ชบ" in [r[0].value for r in again[rio.CODES_SHEET].iter_rows(min_row=2)]
    assert again[month].cell(_row(again[month], "SCR-1"), _col(again[month], day)).value == "ชบ"

    # a CSV by name, with day numbers: the month comes from the request
    malee = next(p["name"] for p in admin.staff()["staff"] if p["code"] == "SCR-2")
    csv_text = f"ชื่อ,{day.day}\n{malee},OC\n"
    report = rio.import_roster(admin, "roster.csv", csv_text.encode("utf-8-sig"), commit=True, month=month)
    assert report["committed"] and report["from"] == day.isoformat()
    assert _shifts(db_url, "SCR-2", day) == [("ON_CALL", "16:00", "08:00", None)]
