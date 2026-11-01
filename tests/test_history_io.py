"""Past surgeries imported from Excel for the duration forecast (and the before/after report).

Cell parsing needs no database; importing runs against a real PostgreSQL server (TEST_DATABASE_URL).
"""
from __future__ import annotations

import io
from datetime import date, datetime, time, timedelta

import psycopg
import pytest

from app import history_io as hx
from app.store import StoreError
from test_app import OWNER, fresh, make_user, owner, session  # noqa: F401  (fixtures)
from test_db import ADMIN_URL, _rows, db_url, store  # noqa: F401  (fixtures)

needs_db = pytest.mark.skipif(not ADMIN_URL, reason="set TEST_DATABASE_URL to run database tests")
ADMIN = {"name": "Admin A", "username": "a", "role": "admin", "staff_id": None, "user_id": "u-admin"}


# ------------------------------------------------------------------ cells

def test_dates_in_thai_and_international_forms():
    for text in ("31/10/2569", "31/10/2026", "2026-10-31", "31-10-69", "31.10.2569", datetime(2026, 10, 31, 8, 0),
                 date(2026, 10, 31), "31/10/2569 08:30"):
        assert hx.parse_date(text) == date(2026, 10, 31), text
    assert hx.parse_date(46326) == date(1899, 12, 30) + timedelta(days=46326)  # an Excel day number
    for bad in ("31/13/2569", "yesterday", "2026/31/10"):
        with pytest.raises(ValueError, match="วันที่"):
            hx.parse_date(bad)


def test_times():
    for value in ("08:30", "8.30", "0830", "08:30:00", "8:30 น.", time(8, 30), datetime(2026, 1, 1, 8, 30), 510 / 1440):
        assert hx.parse_clock(value) == 510, value
    for bad in ("25:00", "8", "late"):
        with pytest.raises(ValueError, match="เวลา"):
            hx.parse_clock(bad)


def test_headers_are_found_and_patient_columns_left_out():
    grid = [["รายงานผ่าตัด ต.ค. 69"], [], ["HN", "ชื่อผู้ป่วย", "Date", "Surgeon", "Operation", "In", "Out", "OR"],
            ["123", "นาย ก", "31/10/2569", "SUR-GEN", "x", "08:00", "09:00", "OR-1"]]
    head, cols, patient = hx._table(grid)
    assert head == 2 and cols == {"date": 2, "surgeon": 3, "procedure": 4, "in": 5, "out": 6, "room": 7}
    assert patient == ["HN", "ชื่อผู้ป่วย"]
    assert hx._table([["a", "b"], ["c"]]) is None


# ------------------------------------------------------------------ import

def _file(store, rows: list[list], *, extra_header=None) -> bytes:
    """The downloaded template with ``rows`` filled in."""
    from openpyxl import load_workbook

    from app.admin_store import Inventory

    wb = load_workbook(io.BytesIO(hx.template(Inventory(store))))
    ws = wb[hx.SHEET]
    if extra_header:
        ws.cell(1, 10, extra_header)
    for r in rows:
        ws.append(r)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _thai(day: date) -> str:
    return f"{day.day:02d}/{day.month:02d}/{day.year + 543}"


@needs_db
def test_import_checks_then_saves_and_the_forecast_learns(store, db_url):
    from app.admin_store import Inventory

    admin = Inventory(store)
    with psycopg.connect(db_url, autocommit=True) as conn:
        conn.execute("UPDATE procedures SET icd9_cm_code = '51.23' WHERE name = 'Laparoscopic cholecystectomy'")
    today = store.db.today()
    d1, d2 = today - timedelta(days=30), today - timedelta(days=60)
    rows = [
        [_thai(d1), "Dr.Napat", "Laparoscopic cholecystectomy", "08:00", "09:40", None, 90, "OR-1", "นัดล่วงหน้า", "HN-1"],
        [d2.isoformat(), "SUR-GEN", "51.23", None, None, 95, None, None, None, "HN-2"],          # staff code, ICD-9 code
        [_thai(d2), "นพ. Napat", "laparoscopic  cholecystectomy", "23:30", "01:05", None, None, "OR-2", "ฉุกเฉิน"],  # overnight
        [_thai(d1), "Dr.Napat", "Laparoscopic cholecystectomy", "08:00", "09:40", None, 90, "OR-1", "นัดล่วงหน้า"],  # same case
        [_thai(d1), "Dr.Nobody", "Laparoscopic cholecystectomy", "08:00", "09:00"],                # unknown surgeon
        [_thai(today + timedelta(days=3)), "Dr.Napat", "Laparoscopic cholecystectomy", None, None, 90],  # future
        [_thai(d1), "Dr.Napat", "Laparoscopic cholecystectomy"],                                     # no time
        [_thai(today - timedelta(days=900)), "Dr.Napat", "Laparoscopic cholecystectomy", None, None, 90],  # > 2 years
        [_thai(d1), "Dr.Napat", "Laparoscopic cholecystectomy", None, None, 900],                   # outlier
        [_thai(d1), "Dr.Napat", "Laparoscopic cholecystectomy", None, None, 100, None, None, "เร็วๆ"],  # bad urgency
    ]
    content = _file(store, rows, extra_header="HN")
    with store._transaction(write=False) as conn:
        before = len(store.db.history(conn))
    check = hx.import_history(admin, "history.xlsx", content, commit=False, actor=ADMIN)
    assert (check["rows"], check["ready"], check["error_count"], check["old"], check["outlier_count"], check["duplicates"]) == (
        10, 3, 4, 1, 1, 1)
    messages = " | ".join(e["message"] for e in check["errors"])
    for part in ("ไม่พบศัลยแพทย์ Dr.Nobody", "ยังมาไม่ถึง", "ใส่เวลาเข้าห้องและออกห้อง", "ความเร่งด่วนไม่ถูกต้อง"):
        assert part in messages
    assert any("ไม่นำเข้าคอลัมน์ข้อมูลผู้ป่วย" in w for w in check["warnings"])
    assert check["pairs"] == [{"surgeon": "Dr.Napat (General)", "procedure": "Laparoscopic cholecystectomy", "n": 3, "median": 95}]
    assert _rows(db_url, "SELECT count(*) FROM case_history") == [(0,)] and not check["committed"]  # checking saves nothing

    done = hx.import_history(admin, "history.xlsx", content, commit=True, actor=ADMIN)
    assert done["committed"] and done["imported"] == 3
    assert sorted(_rows(db_url, "SELECT actual_mins, urgency, room_code FROM case_history")) == [
        (95, "ELECTIVE", None), (95, "EMERGENCY", "OR-2"), (100, "ELECTIVE", "OR-1")]
    columns = {r[0] for r in _rows(db_url, """SELECT column_name FROM information_schema.columns
                                               WHERE table_name = 'case_history'""")}
    assert not columns & {"patient_hn", "patient_name", "hn"}  # no patient identity is stored
    with store._transaction(write=False) as conn:
        assert len(store.db.history(conn)) == before + 3  # the forecast now learns from them
    card = next(c for c in admin.cards()["cards"] if c["procedure"] == "Laparoscopic cholecystectomy"
                and c["surgeon"].startswith("Dr.Napat"))
    assert card["history"]["n"] >= 3
    assert "นำเข้าประวัติผ่าตัดย้อนหลัง 3 เคส" in _rows(db_url, "SELECT title FROM orchestra_activity ORDER BY created_at DESC LIMIT 1")[0][0]

    again = hx.import_history(admin, "history-copy.xlsx", content, commit=True, actor=ADMIN)
    assert (again["ready"], again["already"], again["imported"], again["committed"]) == (0, 3, 0, False)
    view = hx.overview(admin)
    assert (view["cases"], len(view["batches"]), view["batches"][0]["rows"]) == (3, 1, 3)

    report = store.report((today - timedelta(days=30)).isoformat(), today.isoformat())
    row = next(r for r in report["history"] if r["procedure"] == "Laparoscopic cholecystectomy")
    assert (row["before_n"], row["before_median"]) == (3, 95) and report["history_ready"]

    assert hx.delete_batch(admin, view["batches"][0]["id"], actor=ADMIN) == {"removed": 3}
    assert hx.overview(admin)["cases"] == 0
    with pytest.raises(StoreError, match="ไม่พบชุด"):
        hx.delete_batch(admin, "not-a-batch")


@needs_db
def test_csv_and_other_headers(store):
    from app.admin_store import Inventory

    day = store.db.today() - timedelta(days=10)
    csv = ("HN,Date,Surgeon,Operation,Duration\n"
           f"1,{day.isoformat()},SUR-GEN,Laparoscopic appendectomy,70\n").encode("utf-8-sig")
    out = hx.import_history(Inventory(store), "his-export.csv", csv, commit=False)
    assert out["ready"] == 1 and out["error_count"] == 0
    with pytest.raises(StoreError, match="ไม่พบหัวตาราง"):
        hx.import_history(Inventory(store), "x.csv", b"a,b\n1,2\n", commit=False)
    too_long = ("Date,Surgeon,Operation,Duration\n" + f"{day.isoformat()},SUR-GEN,Lap,70\n" * 5100).encode()
    with pytest.raises(StoreError, match="เกิน 5000 แถว"):
        hx.import_history(Inventory(store), "big.csv", too_long, commit=False)  # not cut off without a word
    named = hx.import_history(Inventory(store), "ประวัติ นายสมชาย ใจดี HN 1234567.csv", csv, commit=False)
    assert "สมชาย" not in named["filename"] and "1234567" not in named["filename"]


def test_bracketed_text_is_removed_in_linear_time():
    import time

    started = time.perf_counter()
    assert hx._norm("Lap. chole (LC)") == "lapchole" and len(hx._person("(" * 60000)) <= 200
    assert time.perf_counter() - started < 0.5


@needs_db
def test_template_lists_surgeons_and_procedures(store):
    from openpyxl import load_workbook

    from app.admin_store import Inventory

    wb = load_workbook(io.BytesIO(hx.template(Inventory(store))))
    assert wb.sheetnames == [hx.GUIDE, hx.SHEET, hx.LISTS] and wb[hx.LISTS].sheet_state == "hidden"
    assert [c.value for c in wb[hx.SHEET][1]][:3] == ["วันที่ผ่าตัด", "ศัลยแพทย์ (รหัสหรือชื่อ)", "หัตถการ (ชื่อหรือรหัส ICD-9)"]
    surgeons = [c.value for c in wb[hx.LISTS]["A"]]
    assert "Dr.Napat (General)" in surgeons
    assert len(wb[hx.SHEET].data_validations.dataValidation) == 3


@needs_db
def test_starting_over_clears_imported_history(store, db_url):
    from app import hospital_io as hio
    from app.admin_store import Inventory

    admin = Inventory(store)
    content = _file(store, [[_thai(store.db.today() - timedelta(days=5)), "SUR-GEN", "Laparoscopic appendectomy",
                             None, None, 60]])
    assert hx.import_history(admin, "h.xlsx", content, commit=True, actor=ADMIN)["imported"] == 1
    assert hio.data_status(admin)["counts"]["ประวัติผ่าตัดย้อนหลัง"] == 1
    hio.wipe_hospital(admin, hio.WIPE_PHRASE, actor=ADMIN)
    assert _rows(db_url, "SELECT (SELECT count(*) FROM case_history), (SELECT count(*) FROM case_history_imports)") == [(0, 0)]


@needs_db
def test_works_before_migration_014(store, db_url):
    from app.admin_store import Inventory

    with psycopg.connect(db_url, autocommit=True) as conn:
        conn.execute("DROP TABLE case_history")
        conn.execute("DROP TABLE case_history_imports")
    admin = Inventory(store)
    assert hx.overview(admin)["ready"] is False
    content = _file(store, [[_thai(store.db.today() - timedelta(days=5)), "SUR-GEN", "Laparoscopic appendectomy",
                             None, None, 60]])
    with pytest.raises(StoreError, match="014"):
        hx.import_history(admin, "h.xlsx", content, commit=False)
    with store._transaction(write=False) as conn:
        assert isinstance(store.db.history(conn), list)  # the forecast still works
    report = store.report()
    assert report["history"] == [] and report["history_ready"] is False


def test_only_admins_import_history(owner):
    assert owner.get("/api/admin/history").status_code == 400  # needs Supabase
    scheduler = make_user(owner, "sched6", "scheduler")
    assert scheduler.get("/api/admin/history").status_code == 403
    assert scheduler.post("/api/admin/history/import", json={"filename": "h.csv", "content": "eA==", "commit": True}).status_code == 403
    assert scheduler.request("DELETE", "/api/admin/history/x", json={}).status_code == 403
